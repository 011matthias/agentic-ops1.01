"""Items 99 + 100 (owner rulings 2026-09-17): a month reads complete only when
it is, and Publish, the month's sign-off, is refused on the route otherwise.

July 2026 read `ready_to_post: true` on 2026-09-17 with 24 purchase charges
holding no receipt and 11 receipts holding no charge: a charge with no receipt
offered nothing to click, so nothing was "undecided". The route published any
run at all, the classic page included, and recorded nobody.

Route-level through the FastAPI app, behind the password gate with a named
operator code so the publisher is a real session label:

* `ready_to_post` keeps its question; `month_complete` and three counts say
  what still blocks the month, and a verdict that exists closes each blocker;
* the publish route refuses an incomplete month with a code, publishes it on
  an explicit override and records the override, and records who and when;
* a classic (statement-first) run never publishes, override or not; a month
  with no statement publishes only on override;
* unpublish says it removed nothing it taught.
"""
from __future__ import annotations

import io
from dataclasses import replace
from datetime import date, datetime
from decimal import Decimal

import pytest

pytest.importorskip("fastapi")
pytest.importorskip("httpx")

from fastapi.testclient import TestClient  # noqa: E402
from openpyxl import Workbook  # noqa: E402
from openpyxl.styles import PatternFill  # noqa: E402

from expense_recon.categorize_charges import derive_subscription_status  # noqa: E402
from expense_recon.llm.client import ExtractedReceipt, MockLLMClient  # noqa: E402
from expense_recon.matching.types import (  # noqa: E402
    Match,
    MatchOutcome,
    MatchType,
    Receipt,
    Transaction,
)
from expense_recon.store import StatementStore  # noqa: E402
from expense_recon.web.app import create_app  # noqa: E402
from expense_recon.web.serialize import snapshot_to_dict  # noqa: E402
from expense_recon.web.store import RunStore  # noqa: E402

JPG = b"\xff\xd8\xff\xe0fake-jpeg-bytes"
CRISS_CODE = "criss-code-1"

HEADERS = ("Date", "Description", "Type", "Amount")
LOVABLE = (datetime(2026, 8, 31), "LOVABLE", "Sale", -15.00)
OBSIDIAN = (datetime(2026, 8, 30), "OBSIDIAN", "Sale", -96.00)
PAYMENT = (datetime(2026, 8, 4), "Payment Thank You-Mobile", "Payment", 111.00)
# The test LLM names a category for an Uber vendor, so the fee carries a guess.
FEE = (datetime(2026, 8, 20), "UBER ONE ANNUAL FEE", "Fee", -95.00)


@pytest.fixture
def client(tmp_path, monkeypatch):
    monkeypatch.setenv("EXPENSE_RECON_RECEIPT_FIRST", "1")
    monkeypatch.delenv("OPENAI_API_KEY", raising=False)
    monkeypatch.delenv("EXPENSE_RECON_OPERATOR_CODE", raising=False)
    monkeypatch.setenv("EXPENSE_RECON_OPERATOR_CODES", f"{CRISS_CODE}:criss")
    app = create_app(tmp_path)
    with TestClient(app) as c:
        c._data_root = tmp_path
        login = c.post("/api/login", json={"code": CRISS_CODE})
        assert login.status_code == 200, login.text
        c.headers.update({"Authorization": f"Bearer {login.json()['token']}"})
        yield c


def _extraction(vendor, total, date, currency="USD"):
    return ExtractedReceipt(
        date=date, total=total, currency=currency, vendor=vendor,
        reference="", line_items=(), confidence=0.9, notes="",
    )


def _wire(monkeypatch, *extractions):
    mock = MockLLMClient(extraction_responses=list(extractions))
    monkeypatch.setattr("expense_recon.cli._build_llm_client", lambda cfg: (mock, None))


def _done(client, resp):
    assert resp.status_code == 200, resp.text
    job = client.get(f"/jobs/{resp.json()['job_id']}").json()
    assert job["status"] == "done", job
    return job


def _create_batch(client, n_receipts, label="August 2026", seed=0, entity=""):
    resp = client.post(
        "/api/expense-batches", data={"legal_entity": entity, "label": label}
    )
    _done(client, resp)
    batch_id = resp.json()["batch_id"]
    if n_receipts:
        files = [
            ("files", (f"r{seed}-{i}.jpg", JPG + bytes([seed, i]),
                       "application/octet-stream"))
            for i in range(n_receipts)
        ]
        _done(client, client.post(f"/api/expense-batches/{batch_id}/receipts", files=files))
    return batch_id


def _attach(client, batch_id, rows, fills=None):
    """`fills` maps a row's Description to a solid fill, the way Criss colours
    her workbook (yellow = booked, gray = booked through recurring)."""
    wb = Workbook()
    ws = wb.active
    ws.append(list(HEADERS))
    for row in rows:
        ws.append(list(row))
        colour = (fills or {}).get(row[1])
        if colour:
            for cell in ws[ws.max_row]:
                cell.fill = PatternFill(
                    start_color=colour, end_color=colour, fill_type="solid",
                )
    buf = io.BytesIO()
    wb.save(buf)
    _done(client, client.post(
        f"/api/expense-batches/{batch_id}/statement",
        files={"statement": (
            "August2026.xlsx", buf.getvalue(),
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )},
        data={
            "account_id": "card-2838",
            "account_legal_entities": '{"card-2838": "Corporate Services"}',
            "account_card_currency": "USD",
        },
    ))


def _attach_csv(client, batch_id, *rows):
    body = "Date,Amount,Vendor\n" + "".join(f"{d},{a},{v}\n" for d, a, v in rows)
    _done(client, client.post(
        f"/api/expense-batches/{batch_id}/statement",
        files={"statement": ("statement.csv", body.encode(), "application/octet-stream")},
        data={
            "account_id": "amex-9001",
            "account_legal_entities": '{"amex-9001": "Corporate Services"}',
            "account_card_currency": "USD",
            "map_transaction_date": "Date",
            "map_amount": "Amount",
            "map_vendor": "Vendor",
        },
    ))


def _view(client, batch_id):
    resp = client.get(f"/api/runs/{batch_id}")
    assert resp.status_code == 200, resp.text
    return resp.json()


def _row(view, vendor):
    return next(r for r in view["rows"] if r["vendor"] == vendor)


def _decide(client, batch_id, tx_id, status):
    resp = client.post(
        f"/api/runs/{batch_id}/decisions",
        json={"transaction_id": tx_id, "status": status},
    )
    assert resp.status_code == 200, resp.text


def _month(client, monkeypatch, rows, *extractions, fills=None):
    _wire(monkeypatch, *extractions)
    batch_id = _create_batch(client, len(extractions))
    _attach(client, batch_id, rows, fills)
    view = _view(client, batch_id)
    # Every clean pairing decided, so `ready_to_post` is true and only the
    # completeness rule can hold the month back.
    for r in view["rows"]:
        if r["turn"] == "decide":
            _decide(client, batch_id, r["transaction_id"], "confirmed")
    return batch_id


# ── item 99: the month reads complete only when it is ─────────────────


def test_a_charge_with_no_receipt_keeps_the_month_incomplete(client, monkeypatch):
    batch = _month(
        client, monkeypatch, [LOVABLE, OBSIDIAN, PAYMENT],
        _extraction("Lovable Labs", "15.00", "2026-08-31"),
    )
    summary = _view(client, batch)["summary"]
    assert summary["n_undecided"] == 0
    assert summary["ready_to_post"] is True, "the old question keeps its answer"
    assert summary["month_complete"] is False
    assert summary["n_charges_need_receipt"] == 1      # OBSIDIAN; the payment is a credit
    assert summary["n_receipts_need_charge"] == 0

    # The verdict that exists for a charge with no receipt: already booked.
    _decide(client, batch, _row(_view(client, batch), "OBSIDIAN")["transaction_id"],
            "already_posted")
    summary = _view(client, batch)["summary"]
    assert summary["n_charges_need_receipt"] == 0
    assert summary["month_complete"] is True


def test_a_receipt_with_no_charge_keeps_the_month_incomplete(client, monkeypatch):
    batch = _month(
        client, monkeypatch, [LOVABLE, PAYMENT],
        _extraction("Lovable Labs", "15.00", "2026-08-31"),
        _extraction("Redis Inc.", "13200.00", "2026-08-12"),
    )
    view = _view(client, batch)
    assert view["summary"]["n_receipts_need_charge"] == 1
    assert view["summary"]["month_complete"] is False

    doc = view["unmatched_receipts"][0]["document_id"]
    resp = client.post(
        f"/api/runs/{batch}/receipts/{doc}/settled-outside", json={"how": "bank_transfer"},
    )
    assert resp.status_code == 200, resp.text
    summary = _view(client, batch)["summary"]
    assert summary["n_receipts_need_charge"] == 0
    assert summary["month_complete"] is True


def test_a_confirmed_private_expense_needs_no_charge(client, monkeypatch):
    """PR #987: a receipt paid on someone's own card will never meet a company
    card charge. Confirmation is the pair (the flag AND who is reimbursed),
    the same rule `_private_reimbursements` applies everywhere else."""
    batch = _month(
        client, monkeypatch, [LOVABLE, PAYMENT],
        _extraction("Lovable Labs", "15.00", "2026-08-31"),
        _extraction("Brauhaus Kuehler Krug", "140.00", "2026-08-12"),
    )
    view = _view(client, batch)
    assert view["summary"]["n_receipts_need_charge"] == 1
    assert view["summary"]["month_complete"] is False
    doc = view["unmatched_receipts"][0]["document_id"]

    # The flag alone is not a confirmation. The routes refuse to write one, so
    # plant it the way an older one-field-at-a-time write left it.
    with RunStore(client._data_root / "recon-web.sqlite") as store:
        store.set_expense_field_override(batch, doc, "private", "1", "2026-09-17T09:00:00")
    assert _view(client, batch)["summary"]["n_receipts_need_charge"] == 1

    resp = client.post(
        f"/api/runs/{batch}/expenses/{doc}/private",
        json={"private": True, "reimburse_to": "Dirk"},
    )
    assert resp.status_code == 200, resp.text
    view = _view(client, batch)
    assert view["summary"]["n_unmatched_rec"] == 1, "the list keeps its question"
    assert view["summary"]["n_receipts_need_charge"] == 0
    assert view["summary"]["month_complete"] is True
    # The gate reads the same rule: a complete month publishes without override.
    resp = client.post(f"/api/runs/{batch}/publish")
    assert resp.status_code == 200, resp.text
    assert resp.json()["published_override"] is False


def test_a_decided_copy_needs_no_charge_until_ruled_not_a_copy(client, monkeypatch):
    """Item 94's shape: a Lovable invoice and its receipt are one purchase.
    The copy the run payload sets aside (`decided_copies`) needs no charge;
    "Not a copy" makes it a receipt waiting for one again."""
    _wire(
        monkeypatch,
        _extraction("Lovable Labs Incorporated", "15.00", "2026-08-31"),
        _extraction("Lovable Labs Incorporated", "15.00", "2026-08-31"),
    )
    resp = client.post(
        "/api/expense-batches",
        data={"legal_entity": "Corporate Services", "label": "August 2026"},
    )
    _done(client, resp)
    batch = resp.json()["batch_id"]
    _done(client, client.post(
        f"/api/expense-batches/{batch}/receipts",
        files=[
            ("files", (name, JPG + str(i).encode(), "application/octet-stream"))
            for i, name in enumerate(["Invoice-HMVWDWIL.jpg", "Receipt-2167-5718.jpg"])
        ],
    ))
    _attach(client, batch, [LOVABLE, PAYMENT])
    for r in _view(client, batch)["rows"]:
        if r["turn"] == "decide":
            _decide(client, batch, r["transaction_id"], "confirmed")

    view = _view(client, batch)
    (copy,) = view["copies_set_aside"]
    assert view["summary"]["n_receipts_need_charge"] == 0
    assert view["summary"]["month_complete"] is True

    (group,) = view["duplicate_groups"]
    resp = client.post(
        f"/api/runs/{batch}/duplicates/resolve",
        json={"group_id": group["group_id"], "resolution": "ignore"},
    )
    assert resp.status_code == 200, resp.text
    view = _view(client, batch)
    assert view["copies_set_aside"] == []
    # Two purchases now and one charge: one of the two documents waits for a
    # charge again. Which one the charge keeps is the matcher's pick between
    # two exact candidates (since item 217 the set-aside copy was the invoice).
    waiting = {r["document_id"] for r in view["unmatched_receipts"]}
    assert len(waiting) == 1
    assert waiting <= {copy["document_id"], copy["duplicate"]["of"]}
    assert view["summary"]["n_receipts_need_charge"] == 1
    assert view["summary"]["month_complete"] is False


def test_a_fee_needs_no_receipt_but_its_guessed_category_is_undecided(client, monkeypatch):
    batch = _month(
        client, monkeypatch, [LOVABLE, FEE, PAYMENT],
        _extraction("Lovable Labs", "15.00", "2026-08-31"),
    )
    view = _view(client, batch)
    fee = _row(view, "UBER ONE ANNUAL FEE")
    assert fee["row_type"] == "fee"
    assert fee["review"]["reason_code"] == "receiptless_suggested", fee["review"]
    summary = view["summary"]
    assert summary["n_charges_need_receipt"] == 0, "a fee line is closed on receipts"
    assert summary["n_charges_category_guessed"] == 1
    assert summary["month_complete"] is False

    _decide(client, batch, fee["transaction_id"], "already_posted")
    summary = _view(client, batch)["summary"]
    assert summary["n_charges_category_guessed"] == 0
    assert summary["month_complete"] is True


def test_a_receipt_another_month_settled_does_not_block_its_own_month(
    client, monkeypatch
):
    """Item 61's live shape: a Google receipt printed 07-31 sits in July, its
    charge posts 08-01 on August's statement, and August settles it. July
    still lists the receipt (with `settled_by`), and it has a charge."""
    _wire(
        monkeypatch,
        _extraction("Google", "71.64", "2026-07-31"),
        _extraction("Lovable", "25.00", "2026-08-15"),
    )
    corp = "Corporate Services"
    july = _create_batch(client, 1, label="July 2026", seed=1, entity=corp)
    august = _create_batch(client, 1, label="August 2026", seed=2, entity=corp)
    _attach_csv(
        client, august,
        ("07/31/2026", "80.28", "OPENAI"),
        ("08/01/2026", "71.64", "GOOGLE WORKSPACE"),
        ("08/15/2026", "25.00", "LOVABLE"),
    )
    google = next(r for r in _view(client, august)["rows"]
                  if r["vendor"] == "GOOGLE WORKSPACE")
    assert google.get("settled_by", {}).get("run_id") == july, google
    _attach_csv(client, july, ("07/10/2026", "12.00", "OBSIDIAN"))

    view = _view(client, july)
    listed = view["unmatched_receipts"]
    assert len(listed) == 1 and listed[0]["settled_by"]["run_id"] == august, listed
    assert view["summary"]["n_unmatched_rec"] == 1
    assert view["summary"]["n_receipts_need_charge"] == 0


# ── owner ruling 2026-09-17: gray is booked through recurring ─────────

GRAY = "FFD9D9D9"
YELLOW = "FFFFEB9C"
NOTION = (datetime(2026, 8, 12), "NOTION LABS", "Sale", -10.00)


def test_a_gray_filled_charge_is_booked_through_recurring_and_closes(
    client, monkeypatch
):
    """July 2026's shape: every charge left without a receipt is gray, which
    Criss's walkthrough calls "já estão no recurring". Gray closes the receipt
    requirement and the guessed category, the way yellow does, and the month
    publishes without an override."""
    batch = _month(
        client, monkeypatch, [LOVABLE, OBSIDIAN, FEE, NOTION, PAYMENT],
        _extraction("Lovable Labs", "15.00", "2026-08-31"),
        fills={OBSIDIAN[1]: GRAY, FEE[1]: GRAY, NOTION[1]: YELLOW},
    )
    view = _view(client, batch)
    obsidian = _row(view, "OBSIDIAN")
    assert obsidian["entry_status"] == "subscription"
    assert obsidian["effective_bucket"] == "unmatched"
    assert "entry_status_source" not in obsidian, "parallel field, absent on a fill"
    fee = _row(view, "UBER ONE ANNUAL FEE")
    assert fee["entry_status"] == "subscription"
    assert fee["review"]["reason_code"] == "receiptless_suggested", fee["review"]
    assert _row(view, "NOTION LABS")["section"] == "posted"

    summary = view["summary"]
    assert summary["n_charges_need_receipt"] == 0
    assert summary["n_charges_category_guessed"] == 0, "its category lives in recurring"
    assert summary["n_charges_closed_recurring"] == 2   # OBSIDIAN + the fee
    assert summary["n_booked_no_receipt"] == 1, "yellow still closes as before (item 102)"
    assert summary["month_complete"] is True

    resp = client.post(f"/api/runs/{batch}/publish")
    assert resp.status_code == 200, resp.text
    assert resp.json()["published_override"] is False


def _snapshot_tx(tid, day, vendor, amount, **kw) -> Transaction:
    return Transaction(
        transaction_id=tid, legal_entity_id="le1", account_id="chase",
        transaction_date=date(2026, 4, day), posting_date=None,
        amount=Decimal(amount), transaction_currency="USD",
        account_card_currency="USD", vendor_from_statement=vendor,
        card_last4="2838", **kw,
    )


def _snapshot(charges, matched_tx):
    receipt = Receipt(
        document_id="m1", legal_entity_id="le1", detected_date=date(2026, 4, 15),
        detected_total=Decimal("20.00"), detected_currency="USD",
        detected_vendor="Cafe",
    )
    outcome = MatchOutcome(
        matches=[Match(
            transaction_id=matched_tx, document_id="m1", match_type=MatchType.EXACT,
            confidence=0.99, reason="exact", score=95,
            amount_score=1.0, date_score=1.0, vendor_score=1.0,
        )],
        unmatched_transactions=[
            c.transaction_id for c in charges if c.transaction_id != matched_tx
        ],
        unmatched_receipts=[],
    )
    return snapshot_to_dict(charges, [receipt], outcome, [])


def _store_run(client, run_id, snapshot):
    with RunStore(client._data_root / "recon-web.sqlite") as db:
        db.create_run(
            run_id=run_id, created_at="2026-05-02T00:00:00", label="April 2026",
            operator=None, summary={}, snapshot=snapshot, config={},
            work_dir=str(client._data_root), llm_enabled=False, has_coa=False,
        )
    return _view(client, run_id)


def test_a_subscription_mark_derived_from_history_closes_nothing(client, tmp_path):
    """`derive_subscription_status` writes the same `entry_status` from vendor
    history. That is the tool's guess that a charge recurs, not Criss's record
    that it was booked, so the charge still needs a receipt. A snapshot saved
    before the provenance existed carries no key and reads as the fill."""
    history = tmp_path / "vendor-history.sqlite"
    with StatementStore(history) as store:
        # ANTHROPIC in February and March: two distinct months before April.
        store.ingest_transactions(
            [
                replace(_snapshot_tx(f"h{m}", 5, "ANTHROPIC", "20.00"),
                        transaction_date=date(2026, m, 5))
                for m in (2, 3)
            ],
            statement_id="hist",
        )
    charges = [
        _snapshot_tx("t_fill", 5, "OBSIDIAN", "96.00", entry_status="subscription"),
        _snapshot_tx("t_derived", 6, "ANTHROPIC", "20.00"),
        _snapshot_tx("t_yellow", 10, "NOTION", "10.00", entry_status="posted"),
        _snapshot_tx("t_matched", 15, "CAFE", "20.00"),
    ]
    with StatementStore(history) as store:
        charges = derive_subscription_status(charges, store)
    by_id = {c.transaction_id: c for c in charges}
    assert by_id["t_derived"].entry_status == "subscription"
    assert by_id["t_derived"].entry_status_source == "derived"
    assert by_id["t_fill"].entry_status_source is None, "a fill is never overwritten"

    view = _store_run(client, "derived", _snapshot(charges, "t_matched"))
    rows = {r["transaction_id"]: r for r in view["rows"]}
    assert rows["t_derived"]["entry_status_source"] == "derived"
    assert "entry_status_source" not in rows["t_fill"]
    summary = view["summary"]
    assert summary["n_charges_need_receipt"] == 1          # t_derived only
    assert summary["n_charges_closed_recurring"] == 1      # t_fill
    assert summary["n_booked_no_receipt"] == 1             # t_yellow
    assert summary["month_complete"] is False

    # A snapshot from before the ruling: no `entry_status_source` key at all.
    old = _snapshot(charges, "t_matched")
    for tx in old["transactions"]:
        tx.pop("entry_status_source", None)
    summary = _store_run(client, "old", old)["summary"]
    assert summary["n_charges_need_receipt"] == 0
    assert summary["n_charges_closed_recurring"] == 2


# ── item 100: the route is the gate ───────────────────────────────────


def test_publish_refuses_an_incomplete_month_and_takes_an_override(client, monkeypatch):
    batch = _month(
        client, monkeypatch, [LOVABLE, OBSIDIAN, PAYMENT],
        _extraction("Lovable Labs", "15.00", "2026-08-31"),
    )
    resp = client.post(f"/api/runs/{batch}/publish")
    assert resp.status_code == 400, resp.text
    body = resp.json()
    assert body["code"] == "month_not_complete"
    assert body["readiness"]["n_charges_need_receipt"] == 1
    assert body["readiness"]["month_complete"] is False
    assert "1 charge still needs a receipt" in body["error"]
    view = _view(client, batch)
    assert view["published"] is False and view["published_by"] is None

    resp = client.post(f"/api/runs/{batch}/publish", json={"override": True})
    assert resp.status_code == 200, resp.text
    assert resp.json()["published_override"] is True
    view = _view(client, batch)
    assert view["published"] is True
    assert view["published_by"] == "criss"
    assert view["published_override"] is True
    assert view["published_at"]


def test_a_complete_month_publishes_and_names_who_published(client, monkeypatch):
    batch = _month(
        client, monkeypatch, [LOVABLE, PAYMENT],
        _extraction("Lovable Labs", "15.00", "2026-08-31"),
    )
    assert _view(client, batch)["summary"]["month_complete"] is True
    resp = client.post(f"/api/runs/{batch}/publish")
    assert resp.status_code == 200, resp.text
    body = resp.json()
    assert body["published"] is True and body["published_by"] == "criss"
    assert body["published_override"] is False
    view = _view(client, batch)
    assert (view["published"], view["published_by"], view["published_override"]) == (
        True, "criss", False,
    )
    state = client.get("/api/operator/state").json()
    assert batch in {r["run_id"] for r in state["published_runs"]}

    resp = client.post(f"/api/runs/{batch}/unpublish")
    assert resp.status_code == 200, resp.text
    memory = resp.json()["memory"]
    assert memory["unlearned"] is False
    assert memory["kept"] is True and memory["trigger"] == "publish"
    view = _view(client, batch)
    assert (view["published"], view["published_by"], view["published_override"]) == (
        False, None, False,
    )


def test_a_month_with_no_statement_publishes_only_on_override(client, monkeypatch):
    _wire(monkeypatch, _extraction("Lovable Labs", "15.00", "2026-08-31"))
    batch = _create_batch(client, 1)
    resp = client.post(f"/api/runs/{batch}/publish")
    assert resp.status_code == 400, resp.text
    assert resp.json()["code"] == "no_statement"
    resp = client.post(f"/api/runs/{batch}/publish", json={"override": True})
    assert resp.status_code == 200, resp.text
    assert resp.json()["published_override"] is True


def test_unpublish_before_any_save_says_nothing_was_kept(client, monkeypatch):
    _wire(monkeypatch, _extraction("Lovable Labs", "15.00", "2026-08-31"))
    batch = _create_batch(client, 1)
    resp = client.post(f"/api/runs/{batch}/unpublish")
    assert resp.status_code == 200, resp.text
    assert resp.json()["memory"] == {"unlearned": False, "kept": False}


def test_a_classic_run_never_publishes(client):
    from pathlib import Path

    examples = Path(__file__).resolve().parent.parent / "examples"
    resp = client.post(
        "/api/runs",
        files={
            "statement": ("statement.example.csv",
                          (examples / "statement.example.csv").read_bytes(), "text/csv"),
            "receipts": ("receipts.example.csv",
                         (examples / "receipts.example.csv").read_bytes(), "text/csv"),
        },
        data={"account_id": "amex-9001", "account_card_currency": "USD"},
    )
    assert resp.status_code == 200, resp.text
    run_id = resp.json().get("run_id") or client.get(
        f"/jobs/{resp.json()['job_id']}"
    ).json()["run_id"]
    for body in (None, {"override": True}):
        resp = client.post(f"/api/runs/{run_id}/publish", json=body)
        assert resp.status_code == 400, resp.text
        assert resp.json()["code"] == "not_a_month"
    assert client.get("/api/operator/state").json()["published_runs"] == []
