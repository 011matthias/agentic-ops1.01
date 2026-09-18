"""Note item T2 (2026-09-18): a statement upload has an identity of its own.

A `statements[]` entry was keyed by `file`, the name on disk, which is made
unique per upload (`statement-2.xlsx`). Two of Criss's per-card exports that
share the bank's filename therefore got two names for what may be one file,
a re-upload of the same workbook got a second name for the same bytes, and
nothing said whether two entries were the same file. `statement_id` is the
sha256 over the stored bytes (16 hex): the same bytes uploaded twice, or
re-read after a restore, yield the same id, and a corrected file yields a
new one. Parallel field, absent (never null) on an entry written before it.

Everything here runs through the FastAPI app: the attach route writes the
entry, the two page GETs read it, the re-read rebuilds it, the writeback
route addresses an upload by it.
"""
from __future__ import annotations

import hashlib
import io
from datetime import datetime
from pathlib import Path

import pytest

pytest.importorskip("fastapi")
pytest.importorskip("httpx")

from fastapi.testclient import TestClient  # noqa: E402
from openpyxl import Workbook, load_workbook  # noqa: E402

from expense_recon.llm.client import ExtractedReceipt, MockLLMClient  # noqa: E402
from expense_recon.web.app import create_app  # noqa: E402
from expense_recon.web.store import RunStore  # noqa: E402

JPG = b"\xff\xd8\xff\xe0fake-jpeg-bytes"

HEADERS = ("Date", "Description", "Type", "Amount")
ROWS = [
    (datetime(2026, 8, 31), "LOVABLE", "Sale", -15.00),
    (datetime(2026, 8, 30), "OBSIDIAN", "Sale", -96.00),
    (datetime(2026, 8, 23), "PRESSMASTER DMCC", "Sale", -135.00),
]
# The same month, one amount corrected: different bytes, different id.
CORRECTED_ROWS = [
    (datetime(2026, 8, 31), "LOVABLE", "Sale", -15.00),
    (datetime(2026, 8, 30), "OBSIDIAN", "Sale", -69.00),
    (datetime(2026, 8, 23), "PRESSMASTER DMCC", "Sale", -135.00),
]
# A second card's export, sharing the bank's filename with the first.
OTHER_CARD_ROWS = [
    (datetime(2026, 8, 12), "UBER TRIP", "Sale", -23.40),
    (datetime(2026, 8, 14), "NOTION LABS", "Sale", -10.00),
]

ATTACH_FORM = {
    "account_id": "card-2838",
    "account_legal_entities": '{"card-2838": "Corporate Services"}',
    "account_card_currency": "USD",
}


@pytest.fixture
def client(tmp_path, monkeypatch):
    monkeypatch.setenv("EXPENSE_RECON_RECEIPT_FIRST", "1")
    monkeypatch.delenv("OPENAI_API_KEY", raising=False)
    mock = MockLLMClient(extraction_responses=[
        ExtractedReceipt(
            date="2026-08-31", total="15.00", currency="USD", vendor="Lovable",
            reference="", line_items=(), confidence=0.9, notes="",
            payment_hint=None,
        ),
    ])
    monkeypatch.setattr(
        "expense_recon.cli._build_llm_client", lambda cfg: (mock, None)
    )
    app = create_app(tmp_path)
    with TestClient(app) as c:
        c._data_root = tmp_path
        yield c


def _done(client, resp):
    assert resp.status_code == 200, resp.text
    job = client.get(f"/jobs/{resp.json()['job_id']}").json()
    assert job["status"] == "done", job
    return job


def _create_batch(client, label="August 2026") -> str:
    resp = client.post(
        "/api/expense-batches", data={"legal_entity": "", "label": label},
    )
    _done(client, resp)
    batch_id = resp.json()["batch_id"]
    _done(client, client.post(
        f"/api/expense-batches/{batch_id}/receipts",
        files=[("files", ("a.jpg", JPG, "application/octet-stream"))],
    ))
    return batch_id


def _xlsx_bytes(rows, headers=HEADERS) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.append(list(headers))
    for row in rows:
        ws.append(list(row))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _attach(client, batch_id, payload: bytes, filename="August2026.xlsx"):
    resp = client.post(
        f"/api/expense-batches/{batch_id}/statement",
        files={"statement": (
            filename, payload,
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )},
        data=dict(ATTACH_FORM),
    )
    return _done(client, resp)


def _expected_id(payload: bytes) -> str:
    return hashlib.sha256(payload).hexdigest()[:16]


def _run(client, batch_id) -> dict:
    resp = client.get(f"/api/runs/{batch_id}")
    assert resp.status_code == 200, resp.text
    return resp.json()


def _grid(client, batch_id) -> dict:
    resp = client.get(f"/api/expense-batches/{batch_id}")
    assert resp.status_code == 200, resp.text
    return resp.json()


def _store(client) -> RunStore:
    return RunStore(Path(client._data_root) / "recon-web.sqlite")


def _stored(client, batch_id) -> dict:
    with _store(client) as store:
        run = store.get_run(batch_id)
        return {
            "statements": list((run.snapshot or {}).get("statements") or []),
            "anchors": dict((run.snapshot or {}).get("statement_anchors") or {}),
            "work_dir": Path(run.work_dir),
        }


def _coverage_row(view: dict, file: str) -> dict:
    rows = [c for c in view["coverage"] if file in c["statements"]]
    assert len(rows) == 1, view["coverage"]
    return rows[0]


# ── the id itself ─────────────────────────────────────────────────────


def test_an_attach_records_the_content_id_on_both_payloads_and_in_the_anchors(client):
    batch_id = _create_batch(client)
    payload = _xlsx_bytes(ROWS)
    _attach(client, batch_id, payload)

    run = _run(client, batch_id)
    grid = _grid(client, batch_id)
    (entry,) = run["statements"]
    sid = entry["statement_id"]
    assert sid == _expected_id(payload), "the id IS the sha256 of the stored bytes"
    assert len(sid) == 16 and all(c in "0123456789abcdef" for c in sid)
    assert grid["statements"] == run["statements"], "one month, one list"

    # `coverage[].statement_ids` rides beside the file names, same type
    # family (a list of strings), on both payloads.
    for view in (run, grid):
        row = _coverage_row(view, "August2026.xlsx")
        assert row["statements"] == ["August2026.xlsx"]
        assert row["statement_ids"] == [sid]

    # The anchors are recorded under the file name AND under the id, so a
    # reader that knows either can find the file's rows.
    stored = _stored(client, batch_id)
    assert stored["anchors"]["August2026.xlsx"], "a workbook has row anchors"
    assert stored["anchors"][sid] == stored["anchors"]["August2026.xlsx"]
    # And the id is NOT machinery: the entry carries it on the page.
    assert "_anchors" not in entry


def test_the_same_bytes_twice_share_one_id_and_a_corrected_file_gets_a_new_one(client):
    batch_id = _create_batch(client)
    payload = _xlsx_bytes(ROWS)
    _attach(client, batch_id, payload)
    _attach(client, batch_id, payload)
    corrected = _xlsx_bytes(CORRECTED_ROWS)
    _attach(client, batch_id, corrected)

    run = _run(client, batch_id)
    first, again, fixed = run["statements"]
    assert first["file"] != again["file"], "two uploads, two names on disk"
    assert first["statement_id"] == again["statement_id"] == _expected_id(payload)
    assert again["n_new"] == 0, "the fold absorbed the re-upload"
    assert fixed["statement_id"] == _expected_id(corrected)
    assert fixed["statement_id"] != first["statement_id"]

    row = _coverage_row(run, first["file"])
    assert row["statements"] == [first["file"], again["file"], fixed["file"]]
    # Two distinct ids for three files: the list says which files are one.
    assert row["statement_ids"] == [first["statement_id"], fixed["statement_id"]]

    stored = _stored(client, batch_id)
    assert stored["anchors"][first["statement_id"]] == stored["anchors"][first["file"]]
    assert stored["anchors"][fixed["statement_id"]] == stored["anchors"][fixed["file"]]


def test_a_reread_keeps_the_id_and_rebuilds_the_id_keyed_anchors(client):
    batch_id = _create_batch(client)
    payload = _xlsx_bytes(ROWS)
    _attach(client, batch_id, payload)
    before = _run(client, batch_id)["statements"]
    anchors_before = _stored(client, batch_id)["anchors"]

    _done(client, client.post(f"/api/expense-batches/{batch_id}/statements/reread"))

    after = _run(client, batch_id)["statements"]
    assert [e["statement_id"] for e in after] == [e["statement_id"] for e in before]
    assert after[0]["file"] == before[0]["file"]
    anchors_after = _stored(client, batch_id)["anchors"]
    sid = after[0]["statement_id"]
    assert set(anchors_after) == {after[0]["file"], sid}
    assert anchors_after[sid] == anchors_after[after[0]["file"]]
    assert anchors_after[sid] == anchors_before[sid], "same bytes, same rows"


def test_an_entry_written_before_the_id_reads_absent_and_gains_one_on_reread(client):
    batch_id = _create_batch(client)
    payload = _xlsx_bytes(ROWS)
    _attach(client, batch_id, payload)
    sid = _expected_id(payload)

    # An entry recorded before 2026-09-18: no `statement_id`, anchors keyed
    # by file only. Written straight into the snapshot, the way the live
    # months hold theirs.
    with _store(client) as store:
        run = store.get_run(batch_id)
        snapshot = dict(run.snapshot or {})
        entries = [dict(e) for e in snapshot.get("statements") or []]
        for e in entries:
            e.pop("statement_id", None)
        snapshot["statements"] = entries
        anchors = dict(snapshot.get("statement_anchors") or {})
        anchors.pop(sid, None)
        snapshot["statement_anchors"] = anchors
        assert store.update_run_snapshot(batch_id, snapshot)

    run = _run(client, batch_id)
    (entry,) = run["statements"]
    assert "statement_id" not in entry, "absent, never null"
    assert _coverage_row(run, "August2026.xlsx")["statement_ids"] == []
    grid = _grid(client, batch_id)
    assert "statement_id" not in grid["statements"][0]

    # The re-read reads the same bytes off disk and records the id.
    _done(client, client.post(f"/api/expense-batches/{batch_id}/statements/reread"))
    run = _run(client, batch_id)
    assert run["statements"][0]["statement_id"] == sid
    assert _coverage_row(run, "August2026.xlsx")["statement_ids"] == [sid]
    assert _stored(client, batch_id)["anchors"][sid]


# ── addressing an upload by id ────────────────────────────────────────


def _writeback_rows(path: Path) -> list[str]:
    wb = load_workbook(path)
    ws = wb.active
    return [str(row[1].value) for row in ws.iter_rows(min_row=2) if row[1].value]


def test_the_writeback_addresses_an_upload_by_id_when_two_share_a_filename(client):
    batch_id = _create_batch(client)
    first = _xlsx_bytes(ROWS)
    other = _xlsx_bytes(OTHER_CARD_ROWS)
    _attach(client, batch_id, first, filename="Chase.xlsx")
    _attach(client, batch_id, other, filename="Chase.xlsx")

    run = _run(client, batch_id)
    entry_first, entry_other = run["statements"]
    assert entry_first["upload_name"] == entry_other["upload_name"] == "Chase.xlsx"
    assert entry_first["file"] != entry_other["file"]
    assert entry_first["statement_id"] != entry_other["statement_id"]

    # Ask for the FIRST export by id. The month's current statement is the
    # second one, so a route that ignored the id and fell back to "current"
    # would annotate the wrong workbook, which is exactly the wrong-cell
    # write the selector exists to prevent.
    work_dir = _stored(client, batch_id)["work_dir"]
    first_written = work_dir / f"{Path(entry_first['file']).stem}-categorized.xlsx"
    other_written = work_dir / f"{Path(entry_other['file']).stem}-categorized.xlsx"
    resp = client.get(
        f"/runs/{batch_id}/statement-categorized.xlsx",
        params={"statement_id": entry_first["statement_id"]},
    )
    assert resp.status_code == 200, resp.text
    assert first_written.is_file(), sorted(p.name for p in work_dir.iterdir())
    assert not other_written.exists(), "the id picked the first export, not the current one"
    assert _writeback_rows(first_written) == ["LOVABLE", "OBSIDIAN", "PRESSMASTER DMCC"]

    # The id decides when both are given.
    resp = client.get(
        f"/runs/{batch_id}/statement-categorized.xlsx",
        params={"file": entry_other["file"],
                "statement_id": entry_first["statement_id"]},
    )
    assert resp.status_code == 200, resp.text
    assert not other_written.exists(), "`?file=` did not override the id"

    # An id the month never recorded is the same 404 an unknown file is.
    resp = client.get(
        f"/runs/{batch_id}/statement-categorized.xlsx",
        params={"statement_id": "0123456789abcdef"},
    )
    assert resp.status_code == 404, resp.text
    assert resp.json()["code"] == "statement_not_workbook"
    assert not other_written.exists(), "an unknown id wrote nothing"

    # `?file=` alone still works exactly as before.
    resp = client.get(
        f"/runs/{batch_id}/statement-categorized.xlsx",
        params={"file": entry_other["file"]},
    )
    assert resp.status_code == 200, resp.text
    assert _writeback_rows(other_written) == ["UBER TRIP", "NOTION LABS"]
