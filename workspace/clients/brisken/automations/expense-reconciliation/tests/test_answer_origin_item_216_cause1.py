"""Item 216 cause 1: every surface reads WHO answered from one mapping.

Reproduced on a copy of July 2026 (backup 2026-09-25 03:05): Criss's own
account on a receiptless charge was stored, reached the grid and both CSVs and
outlived a re-match and the refusal re-run, yet her statement sheet printed it
as `<account> (confirm)`, the tool asking her to confirm her own decision. The
cause was not the write. `ClassificationSource` says HOW an answer was made,
and each surface derived WHO stood behind it on its own: the sheet trusted
LEARNED only, the Zoho journal LEARNED and EDITED, the row review called every
answer but hers "the tool guessed this", the report coloured hers as
needs-review. `answer_origin` is now the one mapping (person | rule |
suggestion), and every one of those surfaces reads it.

Route-level on a GL month with an attached statement: three receiptless
charges, one the model guesses, one the merchant list answers (a rule), one
Criss picks (a person). Each surface is read for each.
"""
from __future__ import annotations

import io
import json
from datetime import datetime

import pytest

pytest.importorskip("fastapi")
pytest.importorskip("httpx")

from fastapi.testclient import TestClient  # noqa: E402
from openpyxl import Workbook, load_workbook  # noqa: E402

from expense_recon.llm.client import ClassificationResult, MockLLMClient  # noqa: E402
from expense_recon.matching.types import (  # noqa: E402
    DECIDED_ORIGINS,
    ORIGIN_PERSON,
    ORIGIN_RULE,
    ORIGIN_SUGGESTION,
    Categorization,
    ClassificationSource,
    answer_origin,
    origin_of_source_value,
)
from expense_recon.web.app import create_app  # noqa: E402
from expense_recon.web.store import RunStore  # noqa: E402
from expense_recon.zoho import curated_leaves  # noqa: E402

CORP = "Corporate Services"
CORP_ORG = "822741658"
CODE = "criss-code-216c1"

HEADERS = ("Date", "Description", "Type", "Amount")
GUESSED = (datetime(2026, 8, 12), "ACME ANALYTICS", "Sale", -41.00)
RULED = (datetime(2026, 8, 14), "FIGMA", "Sale", -45.00)
PICKED = (datetime(2026, 8, 18), "PRESSMASTER FZCO", "Sale", -220.00)
PAYMENT = (datetime(2026, 8, 4), "Payment Thank You-Mobile", "Payment", 306.00)


def _codes() -> tuple[str, str, str]:
    """Three distinct postable leaves of Corporate Services: the model's,
    the merchant list's, and Criss's."""
    codes = sorted(curated_leaves.postable_codes(CORP_ORG))
    return codes[0], codes[1], codes[2]


def _label(code: str) -> str:
    return next(lbl for lbl in curated_leaves.llm_leaf_labels(CORP_ORG)
                if lbl.startswith(code + " "))


def _name(code: str) -> str:
    return curated_leaves.binding(code, CORP_ORG).name


def _chart_accounts(org: str) -> list[dict]:
    from expense_recon.zoho._curated_leaves_data import LEAVES

    return [
        {"account_id": b[org][0], "account_name": b[org][1], "account_code": code,
         "account_type": "expense", "is_active": True, "parent_account_name": None}
        for code, (_branch, b) in LEAVES.items() if org in b
    ]


@pytest.fixture
def web(tmp_path, monkeypatch):
    monkeypatch.setenv("EXPENSE_RECON_RECEIPT_FIRST", "1")
    monkeypatch.delenv("OPENAI_API_KEY", raising=False)
    monkeypatch.delenv("EXPENSE_RECON_CARDS", raising=False)
    monkeypatch.delenv("EXPENSE_RECON_OPERATOR_CODE", raising=False)
    monkeypatch.setenv("EXPENSE_RECON_OPERATOR_CODES", f"{CODE}:criss")
    chart = tmp_path / "coa.json"
    chart.write_text(json.dumps({
        CORP_ORG: {"org": {"name": CORP}, "accounts": _chart_accounts(CORP_ORG)},
    }), encoding="utf-8")
    prov = tmp_path / "coa-provision.json"
    prov.write_text(json.dumps({
        "chart_path": str(chart),
        "entities": {CORP: {"org_id": CORP_ORG}},
    }), encoding="utf-8")
    monkeypatch.setenv("EXPENSE_RECON_COA_PROVISION", str(prov))
    with TestClient(create_app(tmp_path)) as c:
        login = c.post("/api/login", json={"code": CODE})
        assert login.status_code == 200, login.text
        c.headers.update({"Authorization": f"Bearer {login.json()['token']}"})
        yield c


def _done(client, resp):
    assert resp.status_code == 200, resp.text
    job = client.get(f"/jobs/{resp.json()['job_id']}").json()
    assert job["status"] == "done", job
    return job


def _month(web, monkeypatch):
    """A GL month whose statement holds the three charges. The model answers
    every charge it is asked about with `model_code`; the merchant list
    answers FIGMA for Corporate Services with `rule_code`, ahead of it."""
    model_code, rule_code, pick_code = _codes()
    mock = MockLLMClient(
        responses=[ClassificationResult(_label(model_code), None, 0.9, "mock")] * 8)
    monkeypatch.setattr("expense_recon.cli._build_llm_client", lambda cfg: (mock, None))
    r = web.put("/api/settings", json={"merchants": {"Figma": {
        "aliases": ["FIGMA"], "category": rule_code,
        "accounts": {CORP: rule_code}}}})
    assert r.status_code == 200, r.text
    resp = web.post("/api/expense-batches", data={"legal_entity": CORP, "label": "August 2026"})
    _done(web, resp)
    batch = resp.json()["batch_id"]
    wb = Workbook()
    ws = wb.active
    ws.append(list(HEADERS))
    for row in (GUESSED, RULED, PICKED, PAYMENT):
        ws.append(list(row))
    buf = io.BytesIO()
    wb.save(buf)
    _done(web, web.post(
        f"/api/expense-batches/{batch}/statement",
        files={"statement": (
            "August2026.xlsx", buf.getvalue(),
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        data={"account_id": "card-2838",
              "account_legal_entities": json.dumps({"card-2838": CORP}),
              "account_card_currency": "USD"},
    ))
    view = web.get(f"/api/runs/{batch}").json()
    assert view["category_vocabulary"] == "gl", "precondition: a GL month"
    return batch, model_code, rule_code, pick_code


def _rows(web, batch) -> dict:
    return {r["vendor"]: r for r in web.get(f"/api/runs/{batch}").json()["rows"]}


def _sheet_cells(web, batch) -> dict:
    """Description -> the tool's posting-account cell, off the downloaded
    statement workbook."""
    resp = web.get(f"/runs/{batch}/statement-categorized.xlsx")
    assert resp.status_code == 200, resp.text
    ws = load_workbook(io.BytesIO(resp.content)).active
    head = [c.value for c in ws[1]]
    desc = head.index("Description")
    col = next(i for i, v in enumerate(head) if v and "(tool)" in str(v))
    return {row[desc]: row[col] for row in ws.iter_rows(min_row=2, values_only=True)}


def _journal(web, batch) -> str:
    """The Zoho journal, with its receiptless rows switched on (the opt-in
    `zoho.export_receiptless_learned`; off on every live month today, so
    this surface changes nothing Criss downloads until it is turned on)."""
    with RunStore(web.app.state.db_path) as store:
        run = store.get_run(batch)
        config = dict(run.config or {})
        config["zoho"] = {**(config.get("zoho") or {}), "export_receiptless_learned": True}
        assert store.update_run_config(batch, config)
    resp = web.get(f"/runs/{batch}/zoho.csv")
    assert resp.status_code == 200, resp.text
    return resp.text


def test_criss_own_pick_on_a_charge_is_printed_as_hers_on_every_surface(web, monkeypatch):
    batch, model_code, rule_code, pick_code = _month(web, monkeypatch)
    rows = _rows(web, batch)
    # Preconditions: one answer of each kind before she acts.
    assert rows["ACME ANALYTICS"]["charge_category"]["source"] == "VENDOR"
    assert rows["FIGMA"]["charge_category"]["source"] == "REGISTRY"
    assert rows["FIGMA"]["charge_category"]["category"] == rule_code
    assert rows["PRESSMASTER FZCO"]["charge_category"]["source"] == "VENDOR"

    tx = rows["PRESSMASTER FZCO"]["transaction_id"]
    r = web.put(f"/api/runs/{batch}/charges/{tx}/category", json={"category": pick_code})
    assert r.status_code == 200, r.text

    rows = _rows(web, batch)
    picked, ruled, guessed = rows["PRESSMASTER FZCO"], rows["FIGMA"], rows["ACME ANALYTICS"]
    # The view names who answered.
    assert picked["charge_category"]["origin"] == ORIGIN_PERSON
    assert ruled["charge_category"]["origin"] == ORIGIN_RULE
    assert guessed["charge_category"]["origin"] == ORIGIN_SUGGESTION
    assert picked["posting_category"]["origin"] == ORIGIN_PERSON

    # The row asks only about the guess.
    assert picked["review"]["state"] == "none"
    assert ruled["review"]["state"] == "none", ruled["review"]
    assert guessed["review"]["reason_code"] == "receiptless_suggested"

    # Her statement sheet prints her account as hers, and the rule's as the
    # rule's; only the model's guess asks to be confirmed.
    cells = _sheet_cells(web, batch)
    assert cells["PRESSMASTER FZCO"] == _name(pick_code), cells
    assert cells["FIGMA"] == _name(rule_code), cells
    assert cells["ACME ANALYTICS"] == f"{_name(model_code)} (confirm)", cells

    # The Zoho journal books the decided two and never the guess.
    journal = _journal(web, batch)
    assert _name(pick_code) in journal
    assert _name(rule_code) in journal
    assert "ACME ANALYTICS" not in journal


def test_clearing_her_pick_hands_the_row_back_to_the_guess(web, monkeypatch):
    batch, model_code, _rule, pick_code = _month(web, monkeypatch)
    tx = _rows(web, batch)["PRESSMASTER FZCO"]["transaction_id"]
    web.put(f"/api/runs/{batch}/charges/{tx}/category", json={"category": pick_code})
    web.put(f"/api/runs/{batch}/charges/{tx}/category", json={"category": ""})
    row = _rows(web, batch)["PRESSMASTER FZCO"]
    assert row["charge_category"]["origin"] == ORIGIN_SUGGESTION
    assert row["review"]["reason_code"] == "receiptless_suggested"
    assert _sheet_cells(web, batch)["PRESSMASTER FZCO"] == f"{_name(model_code)} (confirm)"
    assert "PRESSMASTER" not in _journal(web, batch)


# ── the one mapping, and that every surface agrees with it ─────────────


@pytest.mark.parametrize(("source", "origin"), [
    (ClassificationSource.EDITED, ORIGIN_PERSON),
    (ClassificationSource.LEARNED, ORIGIN_RULE),
    (ClassificationSource.REGISTRY, ORIGIN_RULE),
    (ClassificationSource.LINE, ORIGIN_SUGGESTION),
    (ClassificationSource.VENDOR, ORIGIN_SUGGESTION),
    (ClassificationSource.REVIEW, None),
    (ClassificationSource.UNCLASSIFIED, None),
])
def test_every_source_has_exactly_one_origin(source, origin):
    assert answer_origin(Categorization(category="X", zoho_account=None, source=source,
                                        confidence=1.0, reasoning="")) == origin
    assert origin_of_source_value(source.value) == origin


def test_no_category_is_no_answer_whatever_the_source():
    cat = Categorization(category=None, zoho_account=None,
                         source=ClassificationSource.EDITED, confidence=1.0, reasoning="")
    assert answer_origin(cat) is None


def test_a_joined_source_answers_its_weakest_line():
    assert origin_of_source_value("EDITED; REGISTRY") == ORIGIN_RULE
    assert origin_of_source_value("EDITED; LINE") == ORIGIN_SUGGESTION
    assert origin_of_source_value("EDITED") == ORIGIN_PERSON
    assert origin_of_source_value("LINE; nonsense") is None
    assert origin_of_source_value("") is None


@pytest.mark.parametrize("source", list(ClassificationSource))
def test_the_offline_surfaces_agree_with_the_mapping(source):
    """The journal, the sheet and the report colour decide from the same
    mapping for every source: a decided answer posts, prints plain and is
    blue; nothing else does."""
    from datetime import date
    from decimal import Decimal

    from expense_recon.matching.types import Transaction
    from expense_recon.output import report_xlsx
    from expense_recon.output.sheet_writeback import _cell_value
    from expense_recon.output.zoho_export import _postable_charge

    cat = Categorization(category="E1", zoho_account="Acct", source=source,
                         confidence=1.0, reasoning="")
    decided = answer_origin(cat) in DECIDED_ORIGINS
    assert _postable_charge(cat) is decided
    tx = Transaction(
        transaction_id="t1", legal_entity_id=CORP, account_id="card-1",
        transaction_date=date(2026, 8, 1), posting_date=None,
        amount=Decimal("-1.00"), transaction_currency="USD",
        account_card_currency="USD", vendor_from_statement="X",
    )
    cell = _cell_value(tx, {}, {}, set(), {"t1"}, None, {"t1": cat})
    assert (cell == "Acct") is decided, cell
    assert (report_xlsx._fill_for_source(source) is report_xlsx.FILL_LEARNED) is decided
