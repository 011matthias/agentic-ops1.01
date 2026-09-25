"""Item 216 Build 2, step 2: the model suggests, a rule or a person decides.

Measured on the 2026-09-25 03:05 backup before the change: on the three GL
months 85 expenses carry only the model's lines (July 27, August 27,
September 31), `expenses.csv` wrote the model's account as the posting on
every one, and 32 of them read `ready`, so "Confirm all Ready" would have
ratified them. Owner ruling 2026-09-25: a model-only line prints as a labelled
suggestion, `suggested: <account>`, which the poster refuses; those rows read
`check`; one click on Confirm (stored `inherited`, never taught) turns the row
back into a posting.

The structure: `is_suggestion_only` beside `answer_origin` names the model's
answer in the Zoho-account vocabulary, and `SUGGESTED_PREFIX` beside the
export's placeholders is the one cell label. Every surface reads those two:
the views (`posting_category` never holds it; `suggested_category` does), the
review (`model_suggestion`, confirmable), `expenses.csv`, the statement sheet,
`report.xlsx`, `reconciled.csv` and the reconciliation PDF. A person's pick
reads `EDITED` in the export (it read LINE, the model's own tier). A
bucket-era month keeps the reading it was booked under.

Pinned through the callers: a receipt upload on a GL month, the confirm
route, the category PUT, and a GL month's statement with a receiptless guess.
"""
from __future__ import annotations

import csv
import io
import json
from datetime import datetime
from decimal import Decimal

import pytest

from expense_recon.matching.types import (
    Categorization,
    ClassificationSource,
    is_suggestion_only,
)
from expense_recon.output.posting_common import (
    SUGGESTED_PREFIX,
    _debit_account_and_note,
    is_suggested_cell,
)
from expense_recon.zoho import curated_leaves

CORP = "Corporate Services"
CORP_ORG = "822741658"
CODE = "criss-code-216s2"


def _leaves(org: str) -> list[str]:
    return sorted(
        c for c in curated_leaves.postable_codes(org)
        if not curated_leaves.has_postable_children(org, c)
    )


LEAF = _leaves(CORP_ORG)[0]
OTHER_LEAF = _leaves(CORP_ORG)[1]


def _name(code: str) -> str:
    return curated_leaves.binding(code, CORP_ORG).name


# ── the predicate and the cell ─────────────────────────────────────────


def test_only_the_models_answer_in_account_vocabulary_is_a_suggestion():
    def cat(source, category=LEAF):
        return Categorization(category, "Some account", 0.9, source, "")

    assert is_suggestion_only(cat(ClassificationSource.LINE))
    assert is_suggestion_only(cat(ClassificationSource.VENDOR))
    for decided in (ClassificationSource.EDITED, ClassificationSource.LEARNED,
                    ClassificationSource.REGISTRY):
        assert not is_suggestion_only(cat(decided))
    # A bucket-era month (never converted) keeps the reading it was booked
    # under: the model's bucket answer still posts there.
    assert not is_suggestion_only(cat(ClassificationSource.LINE, "Travel & Transport"))
    # Only a curated leaf code is the account vocabulary; any other name
    # (an older fixture's free text) is not.
    assert not is_suggestion_only(cat(ClassificationSource.LINE, "Software"))
    assert not is_suggestion_only(cat(ClassificationSource.REVIEW, None))
    assert not is_suggestion_only(None)


def test_the_export_cell_labels_a_suggestion_and_nothing_else():
    line = Categorization(LEAF, "IT Expenses", 0.9, ClassificationSource.LINE, "")
    assert _debit_account_and_note(line, None)[0] == f"{SUGGESTED_PREFIX}IT Expenses"
    rule = Categorization(LEAF, "IT Expenses", 1.0, ClassificationSource.REGISTRY, "")
    assert _debit_account_and_note(rule, None)[0] == "IT Expenses"
    bucket = Categorization("Travel & Transport", "Travel", 0.9, ClassificationSource.LINE, "")
    assert _debit_account_and_note(bucket, None)[0] == "Travel"
    assert is_suggested_cell(f"{SUGGESTED_PREFIX}IT Expenses")
    assert not is_suggested_cell("IT Expenses")


def test_the_poster_and_the_journal_check_refuse_a_suggestion():
    from expense_recon.ingest.chart_of_accounts import Account, ChartOfAccounts
    from expense_recon.zoho import accounts

    chart = ChartOfAccounts(accounts=(
        Account(account_id="1", name="IT Expenses", code=LEAF, account_type="expense",
                is_active=True),
    ))
    ok = accounts.resolve_account_id("IT Expenses", chart)
    assert not isinstance(ok, accounts.AccountRefusal), ok
    refused = accounts.resolve_account_id(f"{SUGGESTED_PREFIX}IT Expenses", chart)
    assert isinstance(refused, accounts.AccountRefusal)
    assert refused.reason == accounts.REASON_PLACEHOLDER


# ── through the callers ────────────────────────────────────────────────

pytest.importorskip("fastapi")
pytest.importorskip("httpx")

from fastapi.testclient import TestClient  # noqa: E402
from openpyxl import Workbook, load_workbook  # noqa: E402

from expense_recon.llm.client import (  # noqa: E402
    ClassificationResult,
    ExtractedLineItem,
    ExtractedReceipt,
    MockLLMClient,
)
from expense_recon.web.app import create_app  # noqa: E402

JPG = b"\xff\xd8\xff\xe0fake-jpeg-bytes"
HEADERS = ("Date", "Description", "Type", "Amount")
GUESSED = (datetime(2026, 8, 12), "ACME ANALYTICS", "Sale", -41.00)
RULED = (datetime(2026, 8, 14), "FIGMA", "Sale", -45.00)
PAYMENT = (datetime(2026, 8, 4), "Payment Thank You-Mobile", "Payment", 86.00)


def _label(code: str) -> str:
    return f"{code} {_name(code)}"


def _chart_accounts(org: str) -> list[dict]:
    from expense_recon.zoho._curated_leaves_data import LEAVES

    return [
        {"account_id": b[org][0], "account_name": b[org][1], "account_code": code,
         "account_type": "expense", "is_active": True, "parent_account_name": None}
        for code, (_branch, b) in LEAVES.items() if org in b
    ]


@pytest.fixture
def web(tmp_path, monkeypatch):
    monkeypatch.setenv("EXPENSE_RECON_RECEIPT_FIRST", "1")
    monkeypatch.delenv("OPENAI_API_KEY", raising=False)
    monkeypatch.delenv("EXPENSE_RECON_CARDS", raising=False)
    monkeypatch.delenv("EXPENSE_RECON_OPERATOR_CODE", raising=False)
    monkeypatch.setenv("EXPENSE_RECON_OPERATOR_CODES", f"{CODE}:criss")
    chart = tmp_path / "coa.json"
    chart.write_text(json.dumps({
        CORP_ORG: {"org": {"name": CORP}, "accounts": _chart_accounts(CORP_ORG)},
    }), encoding="utf-8")
    prov = tmp_path / "coa-provision.json"
    prov.write_text(json.dumps({
        "chart_path": str(chart),
        "entities": {CORP: {"org_id": CORP_ORG}},
    }), encoding="utf-8")
    monkeypatch.setenv("EXPENSE_RECON_COA_PROVISION", str(prov))
    with TestClient(create_app(tmp_path)) as c:
        login = c.post("/api/login", json={"code": CODE})
        assert login.status_code == 200, login.text
        c.headers.update({"Authorization": f"Bearer {login.json()['token']}"})
        yield c


def _done(client, resp):
    assert resp.status_code == 200, resp.text
    job = client.get(f"/jobs/{resp.json()['job_id']}").json()
    assert job["status"] == "done", job
    return job


def _csv_accounts(client, batch: str) -> list[str]:
    text = client.get(f"/runs/{batch}/expenses.csv").text
    return [r["Expense Account"] for r in csv.DictReader(io.StringIO(text))
            if r.get("Expense Date")]


def _receipt_batch(web, monkeypatch) -> tuple[str, dict]:
    mock = MockLLMClient(
        extraction_responses=[ExtractedReceipt(
            date="2026-08-01", total="40.00", currency="USD", vendor="Acme",
            reference="", confidence=0.9, notes="",
            line_items=(ExtractedLineItem("Analytics seat, August", "40.00"),))],
        responses=[ClassificationResult(_label(LEAF), None, 0.9, "mock")],
    )
    monkeypatch.setattr("expense_recon.cli._build_llm_client", lambda cfg: (mock, None))
    resp = web.post("/api/expense-batches", data={"legal_entity": CORP})
    assert resp.status_code == 200, resp.text
    batch = resp.json()["batch_id"]
    _done(web, web.post(
        f"/api/expense-batches/{batch}/receipts",
        files=[("files", ("acme.jpg", JPG, "application/octet-stream"))]))
    view = web.get(f"/api/expense-batches/{batch}").json()
    assert view["category_vocabulary"] == "gl", "precondition: a GL month"
    (expense,) = view["expenses"]
    (line,) = expense["line_items"]
    assert (line["category"], line["source"]) == (LEAF, "LINE"), "precondition: the model's line"
    return batch, expense


def test_a_model_line_is_a_suggestion_until_a_person_confirms_it(web, monkeypatch):
    batch, expense = _receipt_batch(web, monkeypatch)

    # The grid: no posting, a labelled suggestion, a question, a Confirm.
    assert expense["posting_category"] is None, expense["posting_category"]
    sugg = expense["suggested_category"]
    assert (sugg["category"], sugg["origin"], sugg["source"]) == (LEAF, "suggestion", "llm")
    assert sugg["zoho_account"] == _name(LEAF)
    assert expense["review"]["state"] == "check", expense["review"]
    assert expense["review"]["reason_code"] == "model_suggestion"
    assert expense["category_confirmable"] is True
    (part,) = expense["books_as"]
    assert part == {"account": f"{SUGGESTED_PREFIX}{_name(LEAF)}", "unassigned": False,
                    "amount": "40.00", "suggested": True}, part
    # The poster's input carries the label, never the account.
    assert _csv_accounts(web, batch) == [f"{SUGGESTED_PREFIX}{_name(LEAF)}"]

    # One click on Confirm: stored `inherited`, and the row posts.
    doc = expense["document_id"]
    r = web.post(f"/api/runs/{batch}/expenses/{doc}/confirm-category")
    assert r.status_code == 200, r.text
    (after,) = web.get(f"/api/expense-batches/{batch}").json()["expenses"]
    assert after["posting_category"]["origin"] == "person", after["posting_category"]
    assert after["posting_category"]["zoho_account"] == _name(LEAF)
    assert "suggested_category" not in after
    assert after["review"]["reason_code"] != "model_suggestion", after["review"]
    assert after["category_confirmable"] is False
    assert "suggested" not in after["books_as"][0]
    assert _csv_accounts(web, batch) == [_name(LEAF)]


def test_a_persons_pick_posts_as_theirs_not_as_the_models(web, monkeypatch):
    """The category PUT: `apply_overrides` read a person's pick as LINE, the
    model's own tier, so the export would have labelled it a suggestion."""
    batch, expense = _receipt_batch(web, monkeypatch)
    r = web.put(f"/api/runs/{batch}/expenses/{expense['document_id']}",
                json={"field": "category", "value": OTHER_LEAF})
    assert r.status_code == 200, r.text
    (after,) = web.get(f"/api/expense-batches/{batch}").json()["expenses"]
    assert (after["posting_category"]["category"], after["posting_category"]["source"]) == (
        OTHER_LEAF, "override")
    assert "suggested_category" not in after
    assert _csv_accounts(web, batch) == [_name(OTHER_LEAF)]


def _xlsx_strings(content: bytes, sheet: str | None = None) -> list[str]:
    wb = load_workbook(io.BytesIO(content), read_only=True)
    sheets = [wb[sheet]] if sheet is not None else wb.worksheets
    return [v for ws in sheets for row in ws.iter_rows(values_only=True)
            for v in row if isinstance(v, str)]


def test_a_receiptless_guess_is_a_suggestion_on_every_surface_and_a_rule_posts(
    web, monkeypatch,
):
    mock = MockLLMClient(responses=[ClassificationResult(_label(LEAF), None, 0.9, "mock")] * 8)
    monkeypatch.setattr("expense_recon.cli._build_llm_client", lambda cfg: (mock, None))
    r = web.put("/api/settings", json={"merchants": {"Figma": {
        "aliases": ["FIGMA"], "category": OTHER_LEAF, "accounts": {CORP: OTHER_LEAF}}}})
    assert r.status_code == 200, r.text
    resp = web.post("/api/expense-batches", data={"legal_entity": CORP, "label": "August 2026"})
    _done(web, resp)
    batch = resp.json()["batch_id"]
    wb = Workbook()
    ws = wb.active
    ws.append(list(HEADERS))
    for row in (GUESSED, RULED, PAYMENT):
        ws.append(list(row))
    buf = io.BytesIO()
    wb.save(buf)
    _done(web, web.post(
        f"/api/expense-batches/{batch}/statement",
        files={"statement": (
            "August2026.xlsx", buf.getvalue(),
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        data={"account_id": "card-2838",
              "account_legal_entities": json.dumps({"card-2838": CORP}),
              "account_card_currency": "USD"},
    ))
    view = web.get(f"/api/runs/{batch}").json()
    assert view["category_vocabulary"] == "gl", "precondition: a GL month"
    rows = {r["vendor"]: r for r in view["rows"]}

    guessed = rows["ACME ANALYTICS"]
    assert guessed["charge_category"]["origin"] == "suggestion", guessed["charge_category"]
    assert guessed["posting_category"] is None, guessed["posting_category"]
    assert guessed["suggested_category"]["category"] == LEAF
    assert guessed["review"]["reason_code"] == "receiptless_suggested"
    ruled = rows["FIGMA"]
    assert ruled["posting_category"]["origin"] == "rule", ruled["posting_category"]
    assert "suggested_category" not in ruled

    label = f"{SUGGESTED_PREFIX}{_name(LEAF)}"
    rc = list(csv.DictReader(io.StringIO(web.get(f"/runs/{batch}/reconciled.csv").text)))
    charge_cells = [r["Charge posting account"] for r in rc]
    assert label in charge_cells, charge_cells
    assert _name(OTHER_LEAF) in charge_cells, "the rule's account prints as it is"

    report = web.get(f"/runs/{batch}/report.xlsx")
    assert report.status_code == 200, report.text
    # Each sheet on its own: the card tab and the unmatched section both
    # carry the charge, and each writes its own account cell.
    names = load_workbook(io.BytesIO(report.content), read_only=True).sheetnames
    for sheet in ("Unmatched", *[n for n in names if "2838" in n]):
        assert label in _xlsx_strings(report.content, sheet), (sheet, names)

    sheet = web.get(f"/runs/{batch}/statement-categorized.xlsx")
    assert sheet.status_code == 200, sheet.text
    cells = _xlsx_strings(sheet.content)
    assert label in cells, cells
    assert not any(c.endswith("(confirm)") for c in cells), cells
    assert _name(OTHER_LEAF) in cells

    pdf = web.get(f"/runs/{batch}/reconciliation-report.pdf")
    assert pdf.status_code == 200, pdf.text[:200]
    from pypdf import PdfReader

    text = " ".join(
        (p.extract_text() or "") for p in PdfReader(io.BytesIO(pdf.content)).pages
    )
    assert "suggested:" in " ".join(text.split()), "the PDF labels the guess"


def test_a_bucket_era_line_still_posts_and_reads_ready():
    """The months still in the eight-bucket vocabulary (April to June, never
    converted by the owner's order) are untouched: the review reads the LINE
    answer as before."""
    from expense_recon.matching.types import LineItem, Receipt
    from expense_recon.web.service import _matched_category_review

    def rec(category):
        return Receipt(
            document_id="d1", legal_entity_id=CORP, detected_date=None,
            detected_total=Decimal("10"), detected_currency="USD",
            detected_vendor="Acme", line_items=(LineItem(
                description="x", line_total=Decimal("10"), quantity=None,
                unit_price=None, categorization=Categorization(
                    category, "Acct", 0.9, ClassificationSource.LINE, "")),),
        )

    assert _matched_category_review(rec("Travel & Transport"), {})["state"] == "ready"
    gl = _matched_category_review(rec(LEAF), {})
    assert (gl["state"], gl["reason_code"]) == ("check", "model_suggestion")
