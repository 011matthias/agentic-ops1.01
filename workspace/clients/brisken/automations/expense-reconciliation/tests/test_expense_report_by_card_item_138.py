"""Item 138 (owner, 2026-09-17): "the PDFs are not organized by the cards
that were reconciled." This file pins the month expense report half.

A company month no cost center partitions now lists its expenses per card,
in the order and grouping the reconciliation report uses (`card_sections`):

* each card has a caption, a line on what its statement settled (charges,
  matched, still open, booked without a receipt), its table and its sums;
* a receipt a charge holds sits under that charge's card, and when the
  tool resolved the receipt to another card the section says so (item 137);
* an unheld receipt sits under its own resolved card, one with no card
  under "No card", last;
* each card's receipt pages follow that card's sums, before the next card.

A cost-center month and a trip month are untouched, and so is a month whose
receipts reach one section only (tests/test_private_expense.py and
tests/test_copies_out_of_totals.py read that layout off page 1).

Route-level: every assertion reads `GET /runs/{id}/expense-report.pdf` after
the real upload, statement attach and card-pick routes.
"""
from __future__ import annotations

import io
from datetime import datetime

import pytest

pytest.importorskip("fastapi")
pytest.importorskip("httpx")
pytest.importorskip("pypdf")

from fastapi.testclient import TestClient  # noqa: E402
from openpyxl import Workbook  # noqa: E402
from openpyxl.styles import PatternFill  # noqa: E402
from pypdf import PdfReader  # noqa: E402

from expense_recon.web.app import create_app  # noqa: E402
from tests.test_feedback_notes_52_53_62_63 import (  # noqa: E402
    CARDS,
    _done,
    _expense,
    _extraction,
    _pick_card,
    _wire,
)

LABEL_2838 = "Credit Card Chase Visa - 2838"
LABEL_3645 = "Credit Card Chase Visa - 3645"


@pytest.fixture
def client(tmp_path, monkeypatch):
    monkeypatch.setenv("EXPENSE_RECON_RECEIPT_FIRST", "1")
    monkeypatch.delenv("OPENAI_API_KEY", raising=False)
    monkeypatch.delenv("EXPENSE_RECON_CARDS", raising=False)
    with TestClient(create_app(tmp_path)) as c:
        yield c


def _receipt_pdf(tag: str) -> bytes:
    """A real one-page receipt whose text names it, so a page can be traced
    back to the receipt it came from. Long enough that ingest reads the text
    layer rather than rasterizing."""
    from reportlab.pdfgen import canvas

    buf = io.BytesIO()
    c = canvas.Canvas(buf)
    c.drawString(60, 700, f"RECEIPT-{tag} total paid by card, thank you for your order")
    c.showPage()
    c.save()
    return buf.getvalue()


def _month(client, tags: list[str], **data) -> str:
    resp = client.post(
        "/api/expense-batches",
        data={"legal_entity": "", "label": "August 2026", **data},
    )
    _done(client, resp)
    batch = resp.json()["batch_id"]
    files = [
        ("files", (f"{tag.lower()}.pdf", _receipt_pdf(tag), "application/pdf"))
        for tag in tags
    ]
    _done(client, client.post(f"/api/expense-batches/{batch}/receipts", files=files))
    return batch


def _attach(client, batch: str, rows, yellow=()) -> None:
    """The statement, with the rows whose description is in `yellow` filled
    the way Criss marks a charge booked (item 102)."""
    wb = Workbook()
    ws = wb.active
    ws.append(["Card", "Date", "Description", "Type", "Amount"])
    fill = PatternFill(fill_type="solid", fgColor="FFFF00")
    for row in rows:
        ws.append(list(row))
        if row[2] in yellow:
            for cell in ws[ws.max_row]:
                cell.fill = fill
    buf = io.BytesIO()
    wb.save(buf)
    _done(client, client.post(
        f"/api/expense-batches/{batch}/statement",
        files={"statement": (
            "August2026.xlsx", buf.getvalue(),
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )},
        data={
            "account_id": "card-2838",
            "account_legal_entities": '{"card-2838": "Corporate Services"}',
            "account_card_currency": "USD",
        },
    ))


def _pages(client, batch: str) -> list[str]:
    resp = client.get(f"/runs/{batch}/expense-report.pdf")
    assert resp.status_code == 200, resp.text
    return [
        " ".join((p.extract_text() or "").split())
        for p in PdfReader(io.BytesIO(resp.content)).pages
    ]


def _page_of(pages: list[str], needle: str) -> int:
    hits = [i for i, text in enumerate(pages) if needle in text]
    assert len(hits) == 1, (needle, hits, pages)
    return hits[0]


STATEMENT = [
    ("2838", datetime(2026, 8, 30), "OBSIDIAN", "Sale", -96.00),
    ("2838", datetime(2026, 8, 20), "ADOBE", "Sale", -60.00),
    ("3645", datetime(2026, 8, 5), "LOVABLE", "Sale", -25.00),
    ("3645", datetime(2026, 8, 23), "PRESSMASTER DMCC", "Sale", -135.00),
]


def _reconciled_month(client, monkeypatch) -> str:
    """Two registry cards with charges on each. Obsidian prints 2838 and its
    charge holds it; Lovable prints nothing, the only LOVABLE charge is on
    3645, and the reviewer picked 2838 for it, so 3645 holds a receipt whose
    own card is 2838. Staples prints 3645 and nothing holds it; Coffee Bar
    has no card at all. ADOBE is booked yellow with no receipt."""
    client.put("/api/settings", json={"cards": CARDS})
    _wire(
        monkeypatch,
        _extraction("Obsidian", "96.00", "2026-08-30", "Visa ...2838"),
        _extraction("Lovable", "25.00", "2026-08-05"),
        _extraction("Staples", "42.50", "2026-08-10", "Visa ...3645"),
        _extraction("Coffee Bar", "5.00", "2026-08-12"),
    )
    batch = _month(client, ["OBSIDIAN", "LOVABLE", "STAPLES", "COFFEE"])
    _attach(client, batch, STATEMENT, yellow={"ADOBE"})
    lovable = _expense(client, batch, "Lovable")["document_id"]
    resp = _pick_card(client, batch, lovable, "corp-2838")
    assert resp.status_code == 200 and "rematch" in resp.json(), resp.text
    return batch


def test_the_listing_is_sectioned_per_card_in_reconciliation_order(client, monkeypatch):
    batch = _reconciled_month(client, monkeypatch)
    pages = _pages(client, batch)
    text = " ".join(pages)

    # The header total is the whole month's, unchanged by the sectioning.
    assert "4 expenses · USD 168.50" in text
    assert "Listing by card" in text

    # Sections in the view's coverage order, "No card" last, each with its
    # sums line over its own receipts.
    sums_2838 = f"{LABEL_2838}: 1 expense · USD 96.00"
    sums_3645 = f"{LABEL_3645}: 2 expenses · USD 67.50"
    sums_none = "No card: 1 expense · USD 5.00"
    assert text.index(sums_2838) < text.index(sums_3645) < text.index(sums_none)

    # Each card's rows sit under its caption: numbering follows the sections.
    at_2838 = text.index("Listing by card")
    at_3645 = text.index(sums_2838)
    at_none = text.index(sums_3645)
    assert at_2838 < text.index("1 2026-08-30 Obsidian") < at_3645
    assert at_3645 < text.index("2 2026-08-05 Lovable") < at_none
    assert at_3645 < text.index("3 2026-08-10 Staples") < at_none
    assert at_none < text.index("4 2026-08-12 Coffee Bar") < text.index(sums_none)


def test_each_card_says_what_its_statement_settled(client, monkeypatch):
    batch = _reconciled_month(client, monkeypatch)
    view = client.get(f"/api/runs/{batch}").json()
    coverage = {c["label"]: c for c in view["coverage"]}
    # The figures below are the coverage panel's own, read back here so the
    # document is pinned to the screen rather than to a guess.
    assert coverage[LABEL_2838]["n_transactions"] == 2
    assert coverage[LABEL_2838]["n_reconciled"] == 1
    assert coverage[LABEL_2838]["unreconciled_by_ccy"] == {}
    assert coverage[LABEL_3645]["unreconciled_by_ccy"] == {"USD": "135.00"}

    text = " ".join(_pages(client, batch))
    assert (
        "Statement: August2026.xlsx · 2 charges · 1 matched · still open: "
        "nothing · booked without a receipt: 1 charge, USD 60.00"
    ) in text
    assert (
        "Statement: August2026.xlsx · 2 charges · 1 matched · still open: "
        "USD 135.00"
    ) in text
    # 3645 books nothing yellow, so only 2838 carries the booked part.
    assert text.count("booked without a receipt") == 1


def test_a_receipt_held_across_cards_is_named_in_its_charges_section(client, monkeypatch):
    batch = _reconciled_month(client, monkeypatch)
    view = client.get(f"/api/runs/{batch}").json()
    (held,) = [r for r in view["rows"] if r.get("cards_differ")]
    assert "3645" in held["coverage_key"]

    text = " ".join(_pages(client, batch))
    note = (
        "is held on a charge on this card, but its own card is "
        f"{LABEL_2838}."
    )
    assert text.count(note) == 1
    at = text.index(note)
    assert text.index(f"{LABEL_3645}: 2 expenses") < at < text.index("No card: 1 expense")
    assert "Expense 2 (" in text[at - 60:at]
    assert "—" not in text[at - 60:at + len(note)]


def test_each_cards_receipt_pages_follow_its_own_section(client, monkeypatch):
    batch = _reconciled_month(client, monkeypatch)
    pages = _pages(client, batch)

    sec_2838 = _page_of(pages, f"{LABEL_2838}: 1 expense")
    sec_3645 = _page_of(pages, f"{LABEL_3645}: 2 expenses")
    sec_none = _page_of(pages, "No card: 1 expense")
    captions = {
        tag: _page_of(pages, f"Expense {n} · {vendor}")
        for n, tag, vendor in (
            (1, "OBSIDIAN", "Obsidian"), (2, "LOVABLE", "Lovable"),
            (3, "STAPLES", "Staples"), (4, "COFFEE", "Coffee Bar"),
        )
    }
    # Every receipt's own page is the one right behind its caption.
    for tag, page in captions.items():
        assert _page_of(pages, f"RECEIPT-{tag}") == page + 1, (tag, pages)
    # Card A's receipts come before card B's caption, and so on down.
    assert sec_2838 < captions["OBSIDIAN"] < sec_3645
    assert sec_3645 < captions["LOVABLE"] < captions["STAPLES"] < sec_none
    assert sec_none < captions["COFFEE"]
    # A caption owns its page: the next section starts on a page of its own.
    assert "RECEIPT-" not in pages[sec_3645] and "RECEIPT-" not in pages[sec_none]
    # The closing note comes after the last card's receipts and no longer
    # points "below".
    closing = _page_of(pages, "Every amount above is the amount the CSV export writes")
    assert closing > captions["COFFEE"]
    assert "Each card's receipts follow its listing" in pages[closing]


def test_a_statement_less_month_still_sections_and_says_no_statement(client, monkeypatch):
    client.put("/api/settings", json={"cards": CARDS})
    _wire(
        monkeypatch,
        _extraction("Obsidian", "96.00", "2026-08-30", "Visa ...2838"),
        _extraction("Staples", "42.50", "2026-08-10", "Visa ...3645"),
    )
    batch = _month(client, ["OBSIDIAN", "STAPLES"])
    pages = _pages(client, batch)
    text = " ".join(pages)
    assert "Listing by card" in text
    assert text.count("No statement loaded for this card.") == 2
    assert "charges" not in text and "booked without a receipt" not in text
    assert (
        _page_of(pages, f"{LABEL_2838}: 1 expense")
        < _page_of(pages, "Expense 1 · Obsidian")
        < _page_of(pages, f"{LABEL_3645}: 1 expense")
        < _page_of(pages, "Expense 2 · Staples")
    )


def test_a_cost_center_month_keeps_its_partition_and_receipts_at_the_end(client, monkeypatch):
    batch = _reconciled_month(client, monkeypatch)
    client.put("/api/settings", json={
        "cards": CARDS, "cost_centers": {"Lidar": {"kind": "project"}},
    })
    obsidian = _expense(client, batch, "Obsidian")["document_id"]
    resp = client.put(
        f"/api/runs/{batch}/expenses/{obsidian}",
        json={"field": "cost_center", "value": "Lidar"},
    )
    assert resp.status_code == 200, resp.text

    pages = _pages(client, batch)
    text = " ".join(pages)
    assert "Listing by cost center" in text
    assert "Listing by card" not in text
    assert "is held on a charge on this card" not in text
    assert "Statement: August2026.xlsx" not in text
    last_sums = _page_of(pages, "Unassigned: 3 expenses")
    captions = [
        _page_of(pages, f"Expense {n} · ")
        for n in (1, 2, 3, 4)
    ]
    # Today's layout: the listing, then every caption with its pages behind.
    assert last_sums < min(captions)
    assert captions == sorted(captions)
    assert "Each receipt follows behind the expense number it proves." in text


def test_a_trip_keeps_its_per_person_sections(client, monkeypatch):
    client.put("/api/settings", json={"cards": CARDS})
    resp = client.post("/api/trips", json={
        "name": "TEST - Rome", "start": "2026-08-01", "end": "2026-08-31",
        "travelers": ["Dirk Neumann - Corp Services"],
    })
    assert resp.status_code == 200, resp.text
    _wire(
        monkeypatch,
        _extraction("Obsidian", "96.00", "2026-08-30", "Visa ...2838"),
        _extraction("Staples", "42.50", "2026-08-10", "Visa ...3645"),
    )
    files = [
        ("files", (f"{t.lower()}.pdf", _receipt_pdf(t), "application/pdf"))
        for t in ("OBSIDIAN", "STAPLES")
    ]
    resp = client.post("/api/expense-batches", files=files, data={
        "legal_entity": "", "label": "TEST - Rome",
        "batch_type": "trip", "trip_id": resp.json()["trip_id"],
    })
    _done(client, resp)

    text = " ".join(_pages(client, resp.json()["batch_id"]))
    assert "Trip report" in text
    assert "Dirk Neumann - Corp Services: 2 expenses" in text
    assert "Listing by card" not in text
    assert "No statement loaded" not in text
