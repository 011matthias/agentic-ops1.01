"""Item 216 Build 2, step 3: the summary says who answered the categories.

Measured on the 2026-09-25 03:05 backup before the change: the three GL
months carry 164 receiptless charges with the model's guess (July 45, August
89, September 30), and the only count the payloads had,
`n_charges_category_guessed`, read 0 / 0 / 1. That count answers a different
question, whether a guess still BLOCKS the month (items 99 + 100, owner
rulings 2026-09-17): July's 45 sit on charges Criss already booked (25 yellow)
or booked through Zoho recurring (20 gray), and the 76 open ones in August and
September are counted once, under `n_charges_need_receipt`. Nothing said who
answered.

The structure: `summary.categories_by_origin = {person, rule, suggestion,
none}` on both payloads, read off each row's own `posting_category` /
`suggested_category` (built from `answer_origin` / `is_suggestion_only`), so
the count equals the rows. The four sum to `rows[]` on the run view and to
`n_expenses` on the grid. The blocker keeps its meaning; widening it would
have listed the booked and the receipt-owing charges a second time in the
month's readiness pill.

Pinned through the callers: a receipt upload + the confirm route on the grid,
and a statement attach + the charge category PUT on the run view.
"""
from __future__ import annotations

import io
import json
from datetime import datetime

import pytest

pytest.importorskip("fastapi")
pytest.importorskip("httpx")

from fastapi.testclient import TestClient  # noqa: E402
from openpyxl import Workbook  # noqa: E402
from openpyxl.styles import PatternFill  # noqa: E402

from expense_recon.llm.client import (  # noqa: E402
    ClassificationResult,
    ExtractedLineItem,
    ExtractedReceipt,
    MockLLMClient,
)
from expense_recon.web.app import create_app  # noqa: E402
from expense_recon.zoho import curated_leaves  # noqa: E402

CORP = "Corporate Services"
CORP_ORG = "822741658"
CODE = "criss-code-216s3"
JPG = b"\xff\xd8\xff\xe0fake-jpeg-bytes"
HEADERS = ("Date", "Description", "Type", "Amount")
GRAY = "FFD9D9D9"
YELLOW = "FFFFEB9C"
OPEN = (datetime(2026, 8, 12), "ACME ANALYTICS", "Sale", -41.00)
RULED = (datetime(2026, 8, 14), "FIGMA", "Sale", -45.00)
RECURRING = (datetime(2026, 8, 15), "NOTION LABS", "Sale", -10.00)
BOOKED = (datetime(2026, 8, 16), "OBSIDIAN", "Sale", -8.00)
PAYMENT = (datetime(2026, 8, 4), "Payment Thank You-Mobile", "Payment", 104.00)


def _leaves(org: str) -> list[str]:
    return sorted(
        c for c in curated_leaves.postable_codes(org)
        if not curated_leaves.has_postable_children(org, c)
    )


LEAF, OTHER_LEAF = _leaves(CORP_ORG)[:2]


def _label(code: str) -> str:
    return f"{code} {curated_leaves.binding(code, CORP_ORG).name}"


def _chart_accounts(org: str) -> list[dict]:
    from expense_recon.zoho._curated_leaves_data import LEAVES

    return [
        {"account_id": b[org][0], "account_name": b[org][1], "account_code": code,
         "account_type": "expense", "is_active": True, "parent_account_name": None}
        for code, (_branch, b) in LEAVES.items() if org in b
    ]


@pytest.fixture
def web(tmp_path, monkeypatch):
    monkeypatch.setenv("EXPENSE_RECON_RECEIPT_FIRST", "1")
    monkeypatch.delenv("OPENAI_API_KEY", raising=False)
    monkeypatch.delenv("EXPENSE_RECON_CARDS", raising=False)
    monkeypatch.delenv("EXPENSE_RECON_OPERATOR_CODE", raising=False)
    monkeypatch.setenv("EXPENSE_RECON_OPERATOR_CODES", f"{CODE}:criss")
    chart = tmp_path / "coa.json"
    chart.write_text(json.dumps({
        CORP_ORG: {"org": {"name": CORP}, "accounts": _chart_accounts(CORP_ORG)},
    }), encoding="utf-8")
    prov = tmp_path / "coa-provision.json"
    prov.write_text(json.dumps({
        "chart_path": str(chart),
        "entities": {CORP: {"org_id": CORP_ORG}},
    }), encoding="utf-8")
    monkeypatch.setenv("EXPENSE_RECON_COA_PROVISION", str(prov))
    mock = MockLLMClient(
        extraction_responses=[
            ExtractedReceipt(
                date="2026-08-01", total=total, currency="USD", vendor=vendor,
                reference="", confidence=0.9, notes="",
                line_items=(ExtractedLineItem(f"{vendor} seat, August", total),))
            # A Stripe-style invoice + receipt for one purchase (the tool
            # decides the second is a copy), and one more expense.
            for vendor, total in (
                ("Acme", "40.00"), ("Acme", "40.00"), ("Beta Consulting", "900.00"),
            )
        ],
        responses=[ClassificationResult(_label(LEAF), None, 0.9, "mock")] * 8,
    )
    monkeypatch.setattr("expense_recon.cli._build_llm_client", lambda cfg: (mock, None))
    with TestClient(create_app(tmp_path)) as c:
        login = c.post("/api/login", json={"code": CODE})
        assert login.status_code == 200, login.text
        c.headers.update({"Authorization": f"Bearer {login.json()['token']}"})
        yield c


def _done(client, resp):
    assert resp.status_code == 200, resp.text
    job = client.get(f"/jobs/{resp.json()['job_id']}").json()
    assert job["status"] == "done", job
    return job


def _split(person=0, rule=0, suggestion=0, none=0) -> dict:
    return {"person": person, "rule": rule, "suggestion": suggestion, "none": none}


def test_the_grid_counts_a_model_line_as_a_suggestion_until_confirmed(web):
    resp = web.post("/api/expense-batches", data={"legal_entity": CORP})
    assert resp.status_code == 200, resp.text
    batch = resp.json()["batch_id"]
    _done(web, web.post(
        f"/api/expense-batches/{batch}/receipts",
        files=[
            ("files", (name, JPG + str(i).encode(), "application/octet-stream"))
            for i, name in enumerate(
                ["Invoice-HMVWDWIL.jpg", "Receipt-2167-5718.jpg", "beta.jpg"])
        ]))
    view = web.get(f"/api/expense-batches/{batch}").json()
    assert view["category_vocabulary"] == "gl", "precondition: a GL month"
    for e in view["expenses"]:
        assert e["suggested_category"]["origin"] == "suggestion", e
    summary = view["summary"]
    assert (summary["n_copies_set_aside"], summary["n_expenses"]) == (1, 2), summary
    # The decided copy leaves the split with `n_expenses` (item 94).
    assert summary["categories_by_origin"] == _split(suggestion=2), summary
    rows = {
        e["vendor"]["display"]: e for e in view["expenses"]
        if not (e.get("duplicate") or {}).get("is_extra")
    }

    # A bill leaves the split with `n_expenses` (item 218), as it leaves every box.
    beta = rows["Beta Consulting"]["document_id"]
    r = web.put(f"/api/runs/{batch}/expenses/{beta}",
                json={"field": "payment_path", "value": "bill"})
    assert r.status_code == 200, r.text
    summary = web.get(f"/api/expense-batches/{batch}").json()["summary"]
    assert (summary["n_bills"], summary["n_expenses"]) == (1, 1), summary
    assert summary["categories_by_origin"] == _split(suggestion=1), summary

    # One Confirm moves the row from the model's column to the person's.
    acme = rows["Acme"]["document_id"]
    r = web.post(f"/api/runs/{batch}/expenses/{acme}/confirm-category")
    assert r.status_code == 200, r.text
    summary = web.get(f"/api/expense-batches/{batch}").json()["summary"]
    assert summary["categories_by_origin"] == _split(person=1), summary


def _statement(web, batch: str) -> None:
    wb = Workbook()
    ws = wb.active
    ws.append(list(HEADERS))
    fills = {RECURRING[1]: GRAY, BOOKED[1]: YELLOW}
    for row in (OPEN, RULED, RECURRING, BOOKED, PAYMENT):
        ws.append(list(row))
        colour = fills.get(row[1])
        if colour:
            for cell in ws[ws.max_row]:
                cell.fill = PatternFill(
                    start_color=colour, end_color=colour, fill_type="solid",
                )
    buf = io.BytesIO()
    wb.save(buf)
    _done(web, web.post(
        f"/api/expense-batches/{batch}/statement",
        files={"statement": (
            "August2026.xlsx", buf.getvalue(),
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        data={"account_id": "card-2838",
              "account_legal_entities": json.dumps({"card-2838": CORP}),
              "account_card_currency": "USD"},
    ))


def test_the_run_view_counts_every_guess_while_the_blocker_keeps_its_meaning(web):
    r = web.put("/api/settings", json={"merchants": {"Figma": {
        "aliases": ["FIGMA"], "category": OTHER_LEAF, "accounts": {CORP: OTHER_LEAF}}}})
    assert r.status_code == 200, r.text
    resp = web.post("/api/expense-batches", data={"legal_entity": CORP, "label": "August 2026"})
    _done(web, resp)
    batch = resp.json()["batch_id"]
    _statement(web, batch)

    view = web.get(f"/api/runs/{batch}").json()
    assert view["category_vocabulary"] == "gl", "precondition: a GL month"
    rows = {r["vendor"]: r for r in view["rows"]}
    for vendor in (OPEN[1], RECURRING[1], BOOKED[1]):
        assert rows[vendor]["suggested_category"]["origin"] == "suggestion", vendor
    assert rows[RECURRING[1]]["entry_status"] == "subscription", "precondition: gray"
    assert rows[BOOKED[1]]["section"] == "posted", "precondition: yellow"
    assert rows[RULED[1]]["posting_category"]["origin"] == "rule"

    summary = view["summary"]
    # Every guess counted, wherever it sits: open, gray, yellow.
    assert summary["categories_by_origin"] == _split(rule=1, suggestion=3, none=1), summary
    assert sum(summary["categories_by_origin"].values()) == len(view["rows"])
    # The blocker is unchanged: the open guess is counted once, as needing a
    # receipt (beside the ruled charge, which owes one too); gray and yellow
    # guesses block nothing (rulings 2026-09-17).
    assert summary["n_charges_need_receipt"] == 2, summary
    assert summary["n_charges_category_guessed"] == 0, summary

    # A person's pick on the open charge moves one count, in the route's own
    # reply as on the next read.
    tx = rows[OPEN[1]]["transaction_id"]
    r = web.put(f"/api/runs/{batch}/charges/{tx}/category", json={"category": OTHER_LEAF})
    assert r.status_code == 200, r.text
    after = _split(person=1, rule=1, suggestion=2, none=1)
    assert r.json()["summary"]["categories_by_origin"] == after, r.json()["summary"]
    assert web.get(f"/api/runs/{batch}").json()["summary"]["categories_by_origin"] == after
