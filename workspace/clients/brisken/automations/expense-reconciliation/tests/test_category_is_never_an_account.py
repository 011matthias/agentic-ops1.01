"""Item 5: a category never stands in for an account.

Three places wrote `zoho_account or category` into an ACCOUNT column:
`posting_common._debit_account_and_note` (the expense and journal exports,
and through `expense_posting_parts` the grid's "books as") and two cells of
the sheet writeback. It fired whenever no chart was loaded, which is every
entity-less batch, since a multi-entity gate carries no single chart. A
bucket label is no account in any org, and under the GL engine a category is
a leaf code, so either one there reads as an account somebody picked.

The assertions run through the writers (`write_zoho_expense_export`,
`write_sheet_writeback`), not the helper, so a writer that stopped calling it
would go red here.
"""
from __future__ import annotations

import csv
from datetime import date
from decimal import Decimal

from openpyxl import Workbook, load_workbook

from expense_recon.ingest.chart_of_accounts import ChartOfAccounts
from expense_recon.matching.types import (
    Categorization,
    ClassificationSource,
    LineItem,
    Match,
    MatchOutcome,
    MatchType,
    Receipt,
    Transaction,
)
from expense_recon.output.sheet_writeback import write_sheet_writeback
from expense_recon.output.zoho_expense_export import (
    EXPENSE_COLUMNS,
    write_zoho_expense_export,
)

UNMAPPED = "(account unmapped - assign)"
UNCATEGORIZED = "(uncategorized - assign)"
ACCOUNT = "IT: Cloud Subscriptions-Others"

CHART = ChartOfAccounts.from_api(
    [
        {
            "account_id": "4369050000000078239",
            "account_name": ACCOUNT,
            "account_code": "E500010-30",
            "account_type": "expense",
            "is_active": True,
        }
    ]
)


def _line(desc, amount, category, account, source=ClassificationSource.LINE):
    return LineItem(
        description=desc,
        line_total=Decimal(amount),
        categorization=Categorization(
            category=category, zoho_account=account,
            confidence=0.9, source=source, reasoning="t",
        ),
    )


def _receipt(*items, doc_id="r1", total="10.00"):
    return Receipt(
        document_id=doc_id, legal_entity_id="le",
        detected_date=date(2026, 7, 1), detected_total=Decimal(total),
        detected_currency="USD", detected_vendor="Vercel",
        line_items=tuple(items),
    )


def _accounts(tmp_path, receipt, chart=None):
    out = tmp_path / f"export-{'chart' if chart else 'none'}.csv"
    write_zoho_expense_export(
        [receipt], out, chart_of_accounts=chart, default_paid_through="Bank"
    )
    with out.open(newline="", encoding="utf-8") as fh:
        rows = list(csv.reader(fh))
    col = EXPENSE_COLUMNS.index("Expense Account")
    return [r[col] for r in rows[1:]]


# ── the exports ─────────────────────────────────────────────────────


def test_a_bucket_category_with_no_account_is_flagged(tmp_path):
    rec = _receipt(_line("usage", "10.00", "Software & Subscriptions", None))
    assert _accounts(tmp_path, rec) == [UNMAPPED]


def test_a_leaf_code_with_no_account_is_flagged_too(tmp_path):
    """Under the GL engine the category IS a code; written here it would
    read as a chosen account while no account was ever set."""
    rec = _receipt(_line("usage", "10.00", "E500010-30", None))
    assert _accounts(tmp_path, rec) == [UNMAPPED]


def test_the_column_no_longer_depends_on_whether_a_chart_loaded(tmp_path):
    """A charted batch already wrote the marker for this line; an
    uncharted one wrote the category. Same line, same answer now."""
    rec = _receipt(_line("usage", "10.00", "Software & Subscriptions", None))
    assert _accounts(tmp_path, rec) == _accounts(tmp_path, rec, CHART)


def test_a_picked_account_still_passes_through(tmp_path):
    rec = _receipt(_line("usage", "10.00", "E500010-30", ACCOUNT))
    assert _accounts(tmp_path, rec) == [ACCOUNT]
    assert _accounts(tmp_path, rec, CHART) == [ACCOUNT]


def test_a_refused_gl_line_stays_uncategorized(tmp_path):
    """The truthiness gate that must not flip: a refused line carries
    `category=None` and REVIEW, and reads as uncategorized, not unmapped."""
    rec = _receipt(
        _line("usage", "10.00", None, None, source=ClassificationSource.REVIEW)
    )
    assert _accounts(tmp_path, rec) == [UNCATEGORIZED]


# ── the sheet writeback ─────────────────────────────────────────────


def _writeback_cell(tmp_path, receipt):
    src = tmp_path / "chris.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.append(["Date", "Amount", "Description"])
    ws.append(["2026-07-01", 10.0, "VERCEL"])
    wb.save(src)
    tx = Transaction(
        transaction_id="card-1:2", legal_entity_id="le", account_id="card-1",
        transaction_date=date(2026, 7, 1), posting_date=None,
        amount=Decimal("10.00"), transaction_currency="USD",
        account_card_currency="USD", vendor_from_statement="VERCEL",
    )
    outcome = MatchOutcome(
        matches=[Match("card-1:2", receipt.document_id, MatchType.EXACT, 0.99, "x", False)]
    )
    out = write_sheet_writeback(src, tmp_path / "out.xlsx", outcome, [tx], [receipt])
    return load_workbook(out).active.cell(row=2, column=4).value


def test_a_matched_line_with_no_account_writes_the_marker(tmp_path):
    rec = _receipt(_line("usage", "10.00", "Software & Subscriptions", None))
    assert _writeback_cell(tmp_path, rec) == UNMAPPED


def test_a_split_names_its_account_and_flags_the_rest(tmp_path):
    rec = _receipt(
        _line("usage", "6.00", "E500010-30", ACCOUNT),
        _line("support", "4.00", "Software & Subscriptions", None),
    )
    assert _writeback_cell(tmp_path, rec) == f"{ACCOUNT}; {UNMAPPED}"
