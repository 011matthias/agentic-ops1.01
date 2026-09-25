"""A statement row has a type; a card payment is not a refund (backlog item 73).

Note #42, July 2026, on `Payment Thank You-Mobile / Corporate Services`:
"how can this item be 'refund' if you dont even know which card it was payed
with". The row was -9,664.81 USD, August carried the same shape at -7,823.16,
and both are Chase's descriptor for the cardholder paying the card down. The
workbook's own Type column says `Payment`; the parsers read it only as "a
credit", and every credit printed as a refund.

`is_credit` and the matcher are untouched: a payment still leaves the
candidate pool and still sits in the `refund` bucket, which means "money back
to the card, never receipt-matched". What changes is what the tool SAYS the
row is: `rows[].row_type`, off the label the statement printed, with the sign's
old reading as the fallback. Beside it `rows[].entity_source` names where the
row's company came from, so a lent entity reads as lent.

Route-level throughout: workbooks attached through `POST .../statement`, rows
read off `GET /api/runs/{id}`, the two documents downloaded off their routes,
and a snapshot stored before row types existed read back through the view.
"""
from __future__ import annotations

import io
import json
from datetime import date, datetime
from decimal import Decimal

import pytest

pytest.importorskip("fastapi")
pytest.importorskip("httpx")

from fastapi.testclient import TestClient  # noqa: E402
from openpyxl import Workbook, load_workbook  # noqa: E402

from expense_recon.ingest._common import (  # noqa: E402
    CREDIT_ROW_TYPES,
    CREDIT_TYPE_VALUES,
    DEBIT_TYPE_VALUES,
    ROW_TYPES,
    row_type_for_label,
    type_label_from_raw_text,
)
from expense_recon.llm.client import (  # noqa: E402
    ExtractedReceipt,
    FxJudgmentResult,
    MockLLMClient,
)
from expense_recon.matching.types import MatchOutcome, Receipt, Transaction  # noqa: E402
from expense_recon.web.app import create_app  # noqa: E402
from expense_recon.web.serialize import snapshot_to_dict  # noqa: E402
from expense_recon.web.store import RunStore  # noqa: E402

JPG = b"\xff\xd8\xff\xe0fake-jpeg-bytes"

# A Chase activity workbook, printed the way Chase prints it: purchases and
# fees negative, the payment and the return positive. 2838 is in the
# registry, 3645 is not.
CHASE_HEADERS = ("Card", "Date", "Description", "Type", "Amount")
CHASE_ROWS = [
    ("2838", datetime(2026, 8, 31), "LOVABLE", "Sale", -15.00),
    ("2838", datetime(2026, 8, 4), "Payment Thank You-Mobile", "Payment", 7823.16),
    ("2838", datetime(2026, 8, 1), "ANNUAL MEMBERSHIP FEE", "Fee", -150.00),
    ("2838", datetime(2026, 8, 12), "AMAZON MKTPL", "Return", 20.00),
    ("3645", datetime(2026, 8, 30), "OBSIDIAN", "Sale", -96.00),
]

REGISTRY = {
    "corp-2838": {
        "label": "Corporate card (Chase)",
        "digits": ["2838"],
        "entity": "Corporate Services",
    },
}


@pytest.fixture
def client(tmp_path, monkeypatch):
    monkeypatch.setenv("EXPENSE_RECON_RECEIPT_FIRST", "1")
    monkeypatch.delenv("OPENAI_API_KEY", raising=False)
    mock = MockLLMClient(
        extraction_responses=[ExtractedReceipt(
            date="2026-08-31", total="15.00", currency="USD", vendor="LOVABLE",
            reference="", line_items=(), confidence=0.9, notes="",
        )],
        fx_responses=[
            FxJudgmentResult(
                is_match=True, same_purchase_confidence=0.9, implied_rate=1.0,
                converted_amount=Decimal("15.00"), reasoning="same purchase",
            )
        ] * 12,
    )
    monkeypatch.setattr(
        "expense_recon.cli._build_llm_client", lambda cfg: (mock, None)
    )
    app = create_app(tmp_path)
    with TestClient(app) as c:
        c._data_root = tmp_path
        yield c


def _done(client, resp):
    assert resp.status_code == 200, resp.text
    job = client.get(f"/jobs/{resp.json()['job_id']}").json()
    assert job["status"] == "done", job
    return job


def _xlsx_bytes(rows, headers) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.append(list(headers))
    for row in rows:
        ws.append(list(row))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _month(client, rows=CHASE_ROWS, headers=CHASE_HEADERS) -> str:
    """A company month with one receipt and the workbook attached."""
    assert client.put("/api/settings", json={"cards": REGISTRY}).status_code == 200
    resp = client.post(
        "/api/expense-batches", data={"legal_entity": "", "label": "August 2026"},
    )
    _done(client, resp)
    batch_id = resp.json()["batch_id"]
    _done(client, client.post(
        f"/api/expense-batches/{batch_id}/receipts",
        files=[("files", ("a.jpg", JPG, "application/octet-stream"))],
    ))
    _done(client, client.post(
        f"/api/expense-batches/{batch_id}/statement",
        files={"statement": (
            "August2026.xlsx", _xlsx_bytes(rows, headers),
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )},
        data={
            "account_id": "card-2838",
            "account_legal_entities": '{"card-2838": "Corporate Services"}',
            "account_card_currency": "USD",
        },
    ))
    return batch_id


def _rows_by_vendor(client, run_id) -> dict[str, dict]:
    resp = client.get(f"/api/runs/{run_id}")
    assert resp.status_code == 200, resp.text
    return {r["vendor"]: r for r in resp.json()["rows"]}


# ── the vocabulary ───────────────────────────────────────────────────


def test_every_recognised_label_names_a_row_type_of_the_same_direction():
    """A label added to either parser set without a row type fails here, and
    so does a row type whose credit-ness contradicts the sign its label
    canonicalizes to: the partition and the display must never disagree."""
    for label in CREDIT_TYPE_VALUES | DEBIT_TYPE_VALUES:
        row_type = row_type_for_label(label.title())
        assert row_type in ROW_TYPES, label
        assert (row_type in CREDIT_ROW_TYPES) == (label in CREDIT_TYPE_VALUES), label
    assert row_type_for_label("Lastschrift") is None
    assert row_type_for_label("") is None


# ── the named instance, through the route ────────────────────────────


def test_a_card_payment_reads_payment_and_matching_does_not_move(client):
    batch_id = _month(client)
    rows = _rows_by_vendor(client, batch_id)

    assert rows["Payment Thank You-Mobile"]["row_type"] == "payment"
    assert rows["AMAZON MKTPL"]["row_type"] == "refund"
    assert rows["ANNUAL MEMBERSHIP FEE"]["row_type"] == "fee"
    assert rows["LOVABLE"]["row_type"] == "purchase"
    assert rows["OBSIDIAN"]["row_type"] == "purchase"

    # The partition is the old one: both credits sit in the refund bucket,
    # the fee and the purchases do not, and the counts still sum.
    assert rows["Payment Thank You-Mobile"]["amount"] == "-7,823.16"
    assert rows["Payment Thank You-Mobile"]["effective_bucket"] == "refund"
    assert rows["AMAZON MKTPL"]["effective_bucket"] == "refund"
    assert rows["ANNUAL MEMBERSHIP FEE"]["effective_bucket"] != "refund"
    summary = client.get(f"/api/runs/{batch_id}").json()["summary"]
    assert summary["n_refunds"] == 2
    assert summary["invariant_ok"] is True


def test_entity_source_names_where_the_company_came_from(client):
    rows = _rows_by_vendor(client, _month(client))
    # 2838 printed on the row and named by the registry.
    assert rows["Payment Thank You-Mobile"]["legal_entity_id"] == "Corporate Services"
    assert rows["Payment Thank You-Mobile"]["entity_source"] == "card"
    # 3645 printed, not in the registry: blank, and says so.
    assert rows["OBSIDIAN"]["legal_entity_id"] == ""
    assert rows["OBSIDIAN"]["entity_source"] == "none"


def test_a_workbook_with_no_card_column_lends_the_upload_entity(client):
    plain_headers = ("Date", "Description", "Type", "Amount")
    plain_rows = [row[1:] for row in CHASE_ROWS]
    rows = _rows_by_vendor(client, _month(client, plain_rows, plain_headers))
    for row in rows.values():
        assert row["legal_entity_id"] == "Corporate Services", row["vendor"]
        assert row["entity_source"] == "batch", row["vendor"]
    assert rows["Payment Thank You-Mobile"]["row_type"] == "payment"


def test_a_csv_statement_types_its_rows_the_same_way(client):
    csv_text = "Card,Date,Description,Type,Amount\n" + "".join(
        f"{c},{d:%m/%d/%Y},{v},{t},{a:.2f}\n" for c, d, v, t, a in CHASE_ROWS
    )
    assert client.put("/api/settings", json={"cards": REGISTRY}).status_code == 200
    resp = client.post(
        "/api/expense-batches", data={"legal_entity": "", "label": "August 2026"},
    )
    _done(client, resp)
    batch_id = resp.json()["batch_id"]
    _done(client, client.post(
        f"/api/expense-batches/{batch_id}/receipts",
        files=[("files", ("a.jpg", JPG, "application/octet-stream"))],
    ))
    _done(client, client.post(
        f"/api/expense-batches/{batch_id}/statement",
        files={"statement": ("August2026.csv", csv_text.encode(), "text/csv")},
        data={
            "account_id": "card-2838",
            "account_legal_entities": '{"card-2838": "Corporate Services"}',
            "account_card_currency": "USD",
        },
    ))
    rows = _rows_by_vendor(client, batch_id)
    assert rows["Payment Thank You-Mobile"]["row_type"] == "payment"
    assert rows["AMAZON MKTPL"]["row_type"] == "refund"
    assert rows["ANNUAL MEMBERSHIP FEE"]["row_type"] == "fee"


def test_no_type_column_reads_a_credit_by_its_sign(client):
    """The fallback, stated rather than hidden: without a label a credit reads
    by its sign ("refund"), except a card payoff, which the statement names in
    its description (front 5, 2026-09-25: the Chase PDF printed six payoffs as
    refunds). The bucket is the sign's either way."""
    headers = ("Card", "Date", "Description", "Amount")
    rows_in = [(c, d, v, a) for c, d, v, _t, a in CHASE_ROWS]
    rows = _rows_by_vendor(client, _month(client, rows_in, headers))
    assert rows["Payment Thank You-Mobile"]["row_type"] == "payment"
    assert rows["Payment Thank You-Mobile"]["effective_bucket"] == "refund"
    assert rows["AMAZON MKTPL"]["row_type"] == "refund"
    assert rows["LOVABLE"]["row_type"] == "purchase"


# ── the two live months: stored before row types existed ─────────────


def test_a_snapshot_stored_before_row_types_reads_the_printed_label(client):
    """July and August were read on 2026-09-10 and hold no `row_type`. The
    Type cell is still in each row's `raw_text`, so the view reads it back
    without a re-read; a row whose raw text carries no label reads by sign."""
    batch_id = _month(client)
    store = RunStore(client._data_root / "recon-web.sqlite")
    run = store.get_run(batch_id)
    snapshot = json.loads(json.dumps(run.snapshot))
    for tx in snapshot["transactions"]:
        assert tx.pop("row_type", "absent") != "absent"
        if tx["vendor_from_statement"] == "AMAZON MKTPL":
            tx["raw_text"] = "AMAZON MKTPL 20.00"  # a PDF-shaped line
    assert store.update_run_snapshot(batch_id, snapshot)
    store.close()

    rows = _rows_by_vendor(client, batch_id)
    assert rows["Payment Thank You-Mobile"]["row_type"] == "payment"
    assert rows["ANNUAL MEMBERSHIP FEE"]["row_type"] == "fee"
    assert rows["AMAZON MKTPL"]["row_type"] == "refund"


def test_the_raw_text_reader_evaluates_nothing():
    live_shape = (
        "{'Card': 2838, 'Transaction Date': datetime.datetime(2026, 7, 6, 0, 0), "
        "'Post Date': datetime.datetime(2026, 7, 6, 0, 0), 'Description': "
        "'Payment Thank You-Mobile', 'Category': None, 'Type': 'Payment', "
        "'Amount': 9664.81, 'Memo': None}"
    )
    assert type_label_from_raw_text(live_shape) == "Payment"
    assert type_label_from_raw_text("{'Transaction Type': 'Return'}") == "Return"
    assert type_label_from_raw_text("{'Card Type': 'Visa'}") is None
    assert type_label_from_raw_text("__import__('os').getcwd()") is None
    assert type_label_from_raw_text("{'Type': __import__('os').getcwd()}") is None
    assert type_label_from_raw_text("07/06 PAYMENT THANK YOU -9,664.81") is None


# ── month health ─────────────────────────────────────────────────────


def _health_run(client, run_id, row_type):
    tx = Transaction(
        transaction_id="t1", legal_entity_id="le1", account_id="card-2838",
        transaction_date=date(2026, 8, 4), posting_date=None,
        amount=Decimal("-50.00"), transaction_currency="USD",
        account_card_currency="USD", vendor_from_statement="PAYMENT",
        is_credit=True, row_type=row_type,
    )
    rec = Receipt(
        document_id="d1", legal_entity_id="le1", detected_date=date(2026, 8, 4),
        detected_total=Decimal("50.00"), detected_currency="USD",
        detected_vendor="SOMETHING",
    )
    outcome = MatchOutcome(unmatched_receipts=["d1"], refunds=["t1"])
    store = RunStore(client._data_root / "recon-web.sqlite")
    store.create_run(
        run_id=run_id, created_at="2026-09-16T00:00:00", label="health",
        operator=None, summary={}, snapshot=snapshot_to_dict([tx], [rec], outcome, []),
        config={}, work_dir=str(client._data_root), llm_enabled=False, has_coa=False,
    )
    store.close()
    return client.get(f"/api/runs/{run_id}").json()["summary"]["month_health"]


def test_a_card_payment_beside_a_same_amount_receipt_is_not_a_broken_sign(client):
    """`month_health` reads a credit exactly paired with a receipt as a sign
    the workbook never canonicalized. A card payment is a credit by nature;
    it is not evidence of anything. A refund keeps the old reading."""
    payment = _health_run(client, "health-payment", "payment")
    assert payment["state"] == "ok", payment
    assert payment["n_exact_pairs"] == 0

    refund = _health_run(client, "health-refund", "refund")
    assert refund["state"] == "broken", refund
    assert refund["suspects"] == ["sign"]


# ── the documents ────────────────────────────────────────────────────


def test_the_reconciliation_document_calls_the_payment_a_card_payment(client):
    from pypdf import PdfReader

    batch_id = _month(client)
    resp = client.get(f"/runs/{batch_id}/reconciliation-report.pdf")
    assert resp.status_code == 200, resp.text[:300]
    text = " ".join(
        " ".join((p.extract_text() or "") for p in PdfReader(io.BytesIO(resp.content)).pages)
        .split()
    )
    assert "card payment" in text
    assert "refund" in text  # the Return is still one


def test_the_report_workbook_names_the_row_type(client, tmp_path):
    batch_id = _month(client)
    resp = client.get(f"/runs/{batch_id}/report.xlsx")
    assert resp.status_code == 200, resp.text[:300]
    wb = load_workbook(io.BytesIO(resp.content), read_only=True)
    credits: dict[str, str] = {}
    in_credits = False
    for row in wb["Unmatched"].iter_rows(values_only=True):
        first = row[0] if row else None
        if isinstance(first, str) and first.startswith("Refunds / credits"):
            in_credits = True
            continue
        if in_credits:
            if not any(row):
                break
            if first != "Card":  # the section's header row
                credits[row[2]] = row[5]
    assert credits == {
        "Payment Thank You-Mobile": "payment",
        "AMAZON MKTPL": "refund",
    }

    # The CLI's --explain sheet says the same thing in its own casing.
    from expense_recon.output.report_xlsx import write_report
    from expense_recon.web.serialize import snapshot_from_dict

    store = RunStore(client._data_root / "recon-web.sqlite")
    transactions, receipts, outcome, _ = snapshot_from_dict(
        store.get_run(batch_id).snapshot
    )
    store.close()
    out = write_report(outcome, transactions, receipts, tmp_path / "r.xlsx", explain=True)
    explain = {
        row[2]: row[5]
        for row in load_workbook(out, read_only=True)["Explain"].iter_rows(
            min_row=2, values_only=True
        )
        if row and row[2]
    }
    assert explain["Payment Thank You-Mobile"] == "PAYMENT"
    assert explain["AMAZON MKTPL"] == "REFUND"
