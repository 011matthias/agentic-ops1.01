# /// script
# requires-python = ">=3.11"
# dependencies = ["openpyxl>=3.1"]
# ///
"""Roll the Brisken hours tracker into a new month.

`log-brisken-hours.py` appends into the LATEST dated workbook under
`workspace/hours-tracker/`; it never creates one. Every month rollover so far
was a throwaway scratch script (`.scratch/build-july.py`, since deleted), which
is why `feedback_hours_tracker_format` carries a standing warning that a future
rollover must remember to carry the OneAssessment tab forward "else OA hours
have no home". This tool is that seed, made durable.

What a rollover has to get right (all of it learned the hard way):
  - carry EVERY engagement tab forward, including the optional ones;
  - clear the data rows but KEEP their cell styles, because the append tool
    donates row 8's number_format to each new row (a General-formatted donor
    makes Earnings render "28" instead of "28.00 EUR");
  - shrink each Excel Table `ref` to header + one blank buffer row, so the
    month starts genuinely empty and no filler row shows up on a weekly sheet;
  - re-anchor the by-week SUMPRODUCT block (J9:J13) to the new month's Mondays,
    starting at the Monday of the 1st so a day in the month's first partial
    week still lands in a bucket;
  - leave B4 a static month label (the append tool converts it to the live
    MIN/MAX formula on the first add);
  - drop `_meta` and any one-off event tab that carries no engagement table.

Usage:
  uv run tools/roll-hours-month.py --month 2026-09 [--name september]
                                   [--from PATH] [--dry-run] [--force]
"""

import argparse
import calendar
import datetime as dt
import re
import shutil
import sys
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.worksheet.datavalidation import DataValidation

FOLDER = Path("workspace/hours-tracker")
_DATED = re.compile(r"hours-tracker-(\d{4})-(\d{2})")

# Engagement tables, keyed on the stable TABLE name: sheet titles get renamed
# between months while the table name cannot change, because the Overview
# structured refs depend on it.
ENGAGEMENT_TABLES = {"HoursLog", "LeadGenLog", "OneAssessmentLog"}

HEADER_ROW = 7
FIRST_DATA_ROW = 8
LAST_COL = 8               # H
WEEK_ROWS = range(9, 14)   # J9:J13, the by-week block


def latest_book(folder: Path) -> "Path | None":
    cands = []
    for p in folder.glob("hours-tracker-*.xlsx"):
        m = _DATED.search(p.name)
        if m:
            cands.append((m.group(1) + m.group(2), p))
    return max(cands)[1] if cands else None


def mondays_for(year: int, month: int) -> list:
    """Monday of the week holding the 1st, then every later Monday that starts
    a week overlapping the month."""
    first = dt.date(year, month, 1)
    last = dt.date(year, month, calendar.monthrange(year, month)[1])
    monday = first - dt.timedelta(days=first.weekday())
    out = []
    while monday <= last:
        out.append(monday)
        monday += dt.timedelta(days=7)
    return out


def clear_data_rows(ws, table) -> int:
    """Blank A..H below the header, keeping every cell's style. Rows are never
    deleted: the Overview block sits at J1:L13 and delete_rows would eat it."""
    last_ref_row = int(re.sub(r"[A-Z]", "", table.ref.split(":")[1]))
    last_row = max(last_ref_row, ws.max_row)
    cleared = 0
    for r in range(FIRST_DATA_ROW, last_row + 1):
        if any(ws.cell(row=r, column=c).value is not None
               for c in range(1, LAST_COL + 1)):
            cleared += 1
        for c in range(1, LAST_COL + 1):
            ws.cell(row=r, column=c).value = None
    return cleared


def roll(book: Path, dest: Path, year: int, month: int, dry_run: bool) -> int:
    wb = load_workbook(book)

    keep = [ws.title for ws in wb.worksheets
            if any(t in ENGAGEMENT_TABLES for t in ws.tables)]
    if not keep:
        print(f"{book.name}: no engagement tables found", file=sys.stderr)
        return 1
    missing = ENGAGEMENT_TABLES - {t for title in keep for t in wb[title].tables}
    if missing:
        print(f"  note: source carries no {sorted(missing)} tab; not created")

    for title in [ws.title for ws in wb.worksheets if ws.title not in keep]:
        print(f"  drop tab   {title}")
        del wb[title]

    weeks = mondays_for(year, month)
    label = f"{calendar.month_name[month]} {year}"

    for title in keep:
        ws = wb[title]
        tname = next(t for t in ws.tables if t in ENGAGEMENT_TABLES)
        table = ws.tables[tname]
        cleared = clear_data_rows(ws, table)

        table.ref = f"A{HEADER_ROW}:H{FIRST_DATA_ROW}"
        # Re-scope the Billable dropdown onto the buffer row; the append tool
        # widens it again as rows are added.
        ws.data_validations.dataValidation = [
            dv for dv in ws.data_validations.dataValidation
            if dv.formula1 != '"Yes,No"'
        ]
        dv = DataValidation(type="list", formula1='"Yes,No"',
                            allow_blank=True, showDropDown=False)
        ws.add_data_validation(dv)
        dv.add(f"F{FIRST_DATA_ROW}:F{FIRST_DATA_ROW}")

        for i, r in enumerate(WEEK_ROWS):
            ws.cell(row=r, column=10).value = (
                dt.datetime(weeks[i].year, weeks[i].month, weeks[i].day)
                if i < len(weeks) else None
            )
        if len(weeks) > len(WEEK_ROWS):
            print(f"  WARNING: {len(weeks)} weeks touch {label}, only "
                  f"{len(WEEK_ROWS)} by-week rows exist", file=sys.stderr)

        ws["B4"] = label
        ws.freeze_panes = f"A{FIRST_DATA_ROW}"
        print(f"  {title:<24} {tname:<17} cleared {cleared} rows -> {table.ref}")

    print(f"  by-week anchors: {', '.join(str(w) for w in weeks)}")
    print(f"  period label   : {label}")
    if dry_run:
        print(f"DRY RUN: would write {dest}")
        return 0
    wb.save(dest)
    print(f"wrote {dest}")
    return 0


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--month", required=True, help="target month, YYYY-MM")
    ap.add_argument("--name", help="month name in the filename (default: derived)")
    ap.add_argument("--from", dest="src", help="source workbook (default: latest)")
    ap.add_argument("--folder", default=str(FOLDER))
    ap.add_argument("--dry-run", action="store_true")
    ap.add_argument("--force", action="store_true",
                    help="overwrite an existing target month")
    a = ap.parse_args()

    folder = Path(a.folder)
    m = re.fullmatch(r"(\d{4})-(\d{2})", a.month)
    if not m:
        print("--month must be YYYY-MM", file=sys.stderr)
        return 1
    year, month = int(m.group(1)), int(m.group(2))
    if not 1 <= month <= 12:
        print("--month: month out of range", file=sys.stderr)
        return 1
    name = (a.name or calendar.month_name[month]).lower()

    src = Path(a.src) if a.src else latest_book(folder)
    if not src or not src.exists():
        print("no source workbook found", file=sys.stderr)
        return 1
    dest = folder / f"hours-tracker-{a.month}-{name}.xlsx"
    if dest.exists() and not a.force:
        print(f"{dest.name} already exists (use --force to overwrite)",
              file=sys.stderr)
        return 1
    if src.resolve() == dest.resolve():
        print("source and target are the same file", file=sys.stderr)
        return 1

    print(f"roll {src.name} -> {dest.name}")
    if a.dry_run:
        return roll(src, dest, year, month, True)
    shutil.copy2(src, dest)
    try:
        return roll(dest, dest, year, month, False)
    except Exception:
        dest.unlink(missing_ok=True)
        raise


if __name__ == "__main__":
    raise SystemExit(main())
