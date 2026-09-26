# /// script
# requires-python = ">=3.10"
# dependencies = ["openpyxl"]
# ///
"""Add or refresh the Month Summary front tab on a Brisken month workbook.

Matches the shape first used on the August submitted copy: engagement
totals for the month, then a by-week block with hours, earnings and
status. The by-week block covers the whole BILLING CYCLE, not just the
calendar month, so a week that started in the previous month still
appears; August's own summary promised exactly that ("Aug 24 to 31,
follows with the next sheet").

  uv run tools/build-brisken-month-summary.py
  uv run tools/build-brisken-month-summary.py --month 2026-09 --cycle-from 2026-08-24

Derived view. Rewrites only the Month Summary sheet and never touches an
engagement tab. Close the workbook in Excel first.
"""
from __future__ import annotations

import argparse
import datetime as dt
import re
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

ROOT = Path(__file__).resolve().parents[1]
TRACKER = ROOT / "workspace" / "hours-tracker"

ENGAGEMENTS = {
    "HoursLog": "Expense Reconciliation",
    "LeadGenLog": "Lead Generation",
    "OneAssessmentLog": "One Assessment",
}

# Weekly sheets whose send is verified in Sent Items (2026-09-21T20:17Z).
SENT = {
    dt.date(2026, 8, 17): "sheet sent 24 Aug, approved 1 Sep",
    dt.date(2026, 8, 24): "sheet sent 21 Sep",
    dt.date(2026, 8, 31): "sheet sent 21 Sep",
    dt.date(2026, 9, 7): "sheet sent 21 Sep",
    dt.date(2026, 9, 14): "sheet sent 21 Sep",
}

EUR = '#,##0.00" €"'
H2 = "0.00"
INK = "1F2937"
GREY = "6B7280"
RULE = Side(style="thin", color="D1D5DB")
TITLE = "Month Summary"


def span_hours(s, e) -> float:
    a, b = s.hour * 60 + s.minute, e.hour * 60 + e.minute
    return ((b - a) if b >= a else (b - a + 1440)) / 60.0


def read_rows(path: Path):
    wb = load_workbook(path, data_only=False)
    out, rate = [], None
    for ws in wb.worksheets:
        for lo in ws.tables.values():
            if lo.name not in ENGAGEMENTS:
                continue
            if rate is None and isinstance(ws["B5"].value, (int, float)):
                rate = float(ws["B5"].value)
            m = re.match(r"([A-Z]+)(\d+):([A-Z]+)(\d+)", lo.ref)
            for r in range(int(m.group(2)) + 1, int(m.group(4)) + 1):
                d = ws.cell(row=r, column=1).value
                s = ws.cell(row=r, column=3).value
                e = ws.cell(row=r, column=4).value
                bill = ws.cell(row=r, column=6).value
                if not (hasattr(d, "date") and hasattr(s, "hour") and hasattr(e, "hour")):
                    continue
                if str(bill or "Yes").strip().lower() == "no":
                    continue
                out.append((d.date(), ENGAGEMENTS[lo.name], span_hours(s, e)))
    return out, rate


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--month", help="YYYY-MM of the book to stamp (default: newest)")
    ap.add_argument("--cycle-from", default="2026-08-24", help="Monday the uninvoiced cycle starts")
    args = ap.parse_args()
    cycle_from = dt.date.fromisoformat(args.cycle_from)

    books = sorted(TRACKER.glob("hours-tracker-20*.xlsx"))
    target = (
        next(b for b in books if args.month and args.month in b.name) if args.month else books[-1]
    )

    own, rate = read_rows(target)
    rate = rate or 14.0
    cycle = [r for b in books for r in read_rows(b)[0]]

    month_by_eng: dict[str, float] = {}
    for _, eng, h in own:
        month_by_eng[eng] = month_by_eng.get(eng, 0.0) + h
    month_total = sum(month_by_eng.values())
    logged_through = max(d for d, _, _ in own)

    weeks: dict[dt.date, float] = {}
    # Which calendar months a week's rows actually fall in. The Monday alone
    # is not the test: the Aug 31 to Sep 6 week has a Monday in August and
    # every one of its hours in September.
    week_months: dict[dt.date, set[tuple[int, int]]] = {}
    for d, _, h in cycle:
        mon = d - dt.timedelta(days=d.weekday())
        if mon >= cycle_from:
            weeks[mon] = weeks.get(mon, 0.0) + h
            week_months.setdefault(mon, set()).add((d.year, d.month))
    cycle_total = sum(weeks.values())

    wb = load_workbook(target)
    if TITLE in wb.sheetnames:
        del wb[TITLE]
    ws = wb.create_sheet(TITLE, 0)
    ws.sheet_view.showGridLines = False

    def put(r, c, v, bold=False, fmt=None, size=11, color=INK, italic=False):
        cell = ws.cell(row=r, column=c, value=v)
        cell.font = Font(bold=bold, size=size, color=color, italic=italic)
        if fmt:
            cell.number_format = fmt
        return cell

    put(1, 1, "Brisken, Monthly Time Sheet", bold=True, size=14)
    put(3, 1, "Month", bold=True)
    put(3, 2, f"{logged_through:%B %Y} (logged through {logged_through:%b} {logged_through.day})")
    put(4, 1, "Prepared by", bold=True)
    put(4, 2, "Matthias Silva")
    put(5, 1, "Hourly rate", bold=True)
    put(5, 2, rate, fmt=EUR)

    put(7, 1, "Engagement", bold=True)
    put(7, 2, "Hours", bold=True)
    put(7, 3, "Earnings", bold=True)
    for c in (1, 2, 3):
        ws.cell(row=7, column=c).border = Border(bottom=RULE)
    r = 8
    for eng in ENGAGEMENTS.values():
        if eng not in month_by_eng:
            continue
        put(r, 1, eng)
        put(r, 2, round(month_by_eng[eng], 2), fmt=H2)
        put(r, 3, f"=B{r}*$B$5", fmt=EUR)
        r += 1
    put(r, 1, f"{logged_through:%B}" + " total", bold=True)
    put(r, 2, round(month_total, 2), bold=True, fmt=H2)
    put(r, 3, f"=B{r}*$B$5", bold=True, fmt=EUR)
    for c in (1, 2, 3):
        ws.cell(row=r, column=c).border = Border(top=Side(style="medium", color="374151"))

    r += 2
    put(r, 1, "By week", bold=True, size=12)
    r += 1
    for i, h in enumerate(("Week", "Hours", "Earnings", "Status"), start=1):
        cell = put(r, i, h, bold=True)
        cell.border = Border(bottom=RULE)
        if i > 1:
            cell.alignment = Alignment(horizontal="center" if i < 4 else "left")
    r += 1
    today = dt.date.today()
    for mon in sorted(weeks):
        sun = mon + dt.timedelta(days=6)
        label = f"{mon:%b} {mon.day} to " + (f"{sun:%b} {sun.day}" if sun.month != mon.month else str(sun.day))
        put(r, 1, label)
        put(r, 2, round(weeks[mon], 2), fmt=H2)
        put(r, 3, f"=B{r}*$B$5", fmt=EUR)
        if sun >= today:
            status = "in progress, closes " + f"{sun:%b %d}"
        else:
            status = SENT.get(mon, "booked, sheet not yet sent")
        this_month = (logged_through.year, logged_through.month)
        outside = sorted(m for m in week_months[mon] if m != this_month)
        if outside:
            status += "; carried from " + dt.date(outside[0][0], outside[0][1], 1).strftime("%B")
        put(r, 4, status, size=10, color=GREY)
        for c in range(1, 5):
            ws.cell(row=r, column=c).border = Border(bottom=RULE)
        r += 1
    put(r, 1, "Total booked this cycle", bold=True, size=12)
    put(r, 2, round(cycle_total, 2), bold=True, size=12, fmt=H2)
    put(r, 3, f"=B{r}*$B$5", bold=True, size=12, fmt=EUR)
    for c in range(1, 5):
        ws.cell(row=r, column=c).fill = PatternFill("solid", fgColor="F3F4F6")
        ws.cell(row=r, column=c).border = Border(top=Side(style="medium", color="374151"))

    r += 2
    put(
        r,
        1,
        f"The cycle runs from {cycle_from:%d %B %Y}. Earlier weeks were invoiced separately; "
        "invoice 01-01110 covered Aug 3 to 23.",
        size=9,
        italic=True,
        color=GREY,
    )
    put(r + 1, 1, "Entry detail per engagement is on the tabs behind this page.", size=9, italic=True, color=GREY)

    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 12
    ws.column_dimensions["C"].width = 14
    ws.column_dimensions["D"].width = 34

    wb.save(target)

    print(f"{TITLE} written to {target.name}   rate EUR {rate:.2f}/hr\n")
    for eng in ENGAGEMENTS.values():
        if eng in month_by_eng:
            print(f"  {eng:26s} {month_by_eng[eng]:7.2f} h")
    print(f"  {logged_through:%B} total".ljust(28) + f"{month_total:7.2f} h = EUR {month_total * rate:,.2f}\n")
    for mon in sorted(weeks):
        sun = mon + dt.timedelta(days=6)
        print(f"  {mon:%d %b} to {sun:%d %b}   {weeks[mon]:6.2f} h")
    print(f"  {'Total booked this cycle':24s}{cycle_total:7.2f} h = EUR {cycle_total * rate:,.2f}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
