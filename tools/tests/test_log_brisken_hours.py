"""log-brisken-hours.py: workbook resolution and the cross-tab overlap gate.

`--file` exists because the tool otherwise only ever sees the LATEST dated
workbook, so the moment a month rolls over the prior month becomes unwritable;
back-filling 2026-08-24..29 after the September rollover is exactly that case.

The overlap gate is the safety-critical half (owner directive 2026-07-23: a
wall-clock minute is billed once across ALL tabs) and had no test. Both are
driven here on a synthetic tracker, with openpyxl importorskip'd so the suite
still runs in a CI env without the dep.
"""
import datetime as dt
import importlib.util
import subprocess
import sys

import pytest

from hooklib import TOOLS

pytest.importorskip("openpyxl")

SCRIPT = TOOLS / "log-brisken-hours.py"


def _load():
    spec = importlib.util.spec_from_file_location("log_brisken_hours", SCRIPT)
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


lbh = _load()


def _tracker(path, rows=()):
    """A two-tab tracker with the real geometry: header at row 7, data at 8."""
    from openpyxl import Workbook
    from openpyxl.worksheet.table import Table

    wb = Workbook()
    del wb[wb.sheetnames[0]]
    for title, table_name in (("Expense Reconciliation", "HoursLog"),
                              ("Lead Generation", "LeadGenLog")):
        ws = wb.create_sheet(title)
        ws["B5"] = 14
        for col, head in enumerate(["Date", "Task", "Start", "End", "Hours",
                                    "Billable", "Earnings", "Notes"], start=1):
            ws.cell(row=7, column=col, value=head)
        mine = [r for r in rows if r[0] == title]
        for i, (_tab, date, task, start, end) in enumerate(mine):
            r = 8 + i
            ws.cell(row=r, column=1, value=dt.datetime.fromisoformat(date))
            ws.cell(row=r, column=2, value=task)
            ws.cell(row=r, column=3, value=dt.time.fromisoformat(start))
            ws.cell(row=r, column=4, value=dt.time.fromisoformat(end))
            ws.cell(row=r, column=6, value="Yes")
        last = 8 + max(len(mine) - 1, 0)
        ws.tables[table_name] = Table(displayName=table_name, ref=f"A7:H{last}")
    wb.save(path)
    return path


def _run(book, *args):
    return subprocess.run(
        [sys.executable, str(SCRIPT), "--file", str(book), *args],
        capture_output=True, text=True)


# --- --file targets a workbook that is not the latest month ----------------

def test_file_flag_targets_an_older_month(tmp_path):
    older = _tracker(tmp_path / "hours-tracker-2026-08-august.xlsx",
                     [("Expense Reconciliation", "2026-08-24", "older month",
                       "12:30", "13:15")])
    _tracker(tmp_path / "hours-tracker-2026-09-september.xlsx")
    out = _run(older, "--status")
    assert out.returncode == 0, out.stderr
    assert "hours-tracker-2026-08-august.xlsx" in out.stdout
    assert "older month" in out.stdout


def test_missing_file_is_reported_not_silently_redirected(tmp_path):
    out = _run(tmp_path / "nope.xlsx", "--status")
    assert out.returncode == 1
    assert "no dated hours-tracker workbook" in out.stderr


# --- the cross-tab overlap gate --------------------------------------------

def test_overlap_with_an_existing_row_in_another_tab_is_refused(tmp_path):
    book = _tracker(tmp_path / "hours-tracker-2026-09-september.xlsx",
                    [("Expense Reconciliation", "2026-09-08", "recon block",
                      "14:00", "16:00")])
    rows = tmp_path / "rows.json"
    rows.write_text('[{"tab": "lead", "date": "2026-09-08", "start": "15:00",'
                    ' "end": "17:00", "task": "double billed"}]')
    out = _run(book, "--add", str(rows))
    assert out.returncode == 1
    assert "OVERLAP" in out.stderr
    assert "recon block" in out.stderr


def test_touching_endpoints_are_allowed(tmp_path):
    book = _tracker(tmp_path / "hours-tracker-2026-09-september.xlsx",
                    [("Expense Reconciliation", "2026-09-08", "recon block",
                      "14:00", "16:00")])
    rows = tmp_path / "rows.json"
    rows.write_text('[{"tab": "lead", "date": "2026-09-08", "start": "16:00",'
                    ' "end": "17:00", "task": "picks up where it left off"}]')
    out = _run(book, "--add", str(rows), "--dry-run")
    assert out.returncode == 0, out.stderr
    assert "picks up where it left off" in out.stdout


def test_a_row_crossing_midnight_still_blocks_the_next_morning(tmp_path):
    # 22:55-01:45 runs into the next day; a 01:00 row the following morning is
    # the same wall-clock minute and must be refused.
    book = _tracker(tmp_path / "hours-tracker-2026-09-september.xlsx",
                    [("Expense Reconciliation", "2026-09-08", "night block",
                      "22:55", "01:45")])
    rows = tmp_path / "rows.json"
    rows.write_text('[{"tab": "lead", "date": "2026-09-09", "start": "01:00",'
                    ' "end": "02:00", "task": "overlaps the night"}]')
    out = _run(book, "--add", str(rows))
    assert out.returncode == 1
    assert "OVERLAP" in out.stderr


def test_two_new_rows_in_one_batch_cannot_overlap_each_other(tmp_path):
    book = _tracker(tmp_path / "hours-tracker-2026-09-september.xlsx")
    rows = tmp_path / "rows.json"
    rows.write_text(
        '[{"tab": "time", "date": "2026-09-08", "start": "10:00",'
        ' "end": "12:00", "task": "first"},'
        ' {"tab": "lead", "date": "2026-09-08", "start": "11:00",'
        ' "end": "13:00", "task": "second"}]')
    out = _run(book, "--add", str(rows))
    assert out.returncode == 1
    assert "OVERLAP" in out.stderr
